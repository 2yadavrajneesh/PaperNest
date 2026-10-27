import io
import json
import os
import shutil
import tempfile
import unittest
from unittest.mock import patch
import zipfile
from pathlib import Path

os.environ.setdefault("PAPERNEST_TEMP_DIR", str(Path(tempfile.gettempdir()) / "papernest-test-jobs"))

from fastapi.testclient import TestClient
from pypdf import PdfReader
from reportlab.pdfgen import canvas
from PIL import Image
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches
import fitz

from app.main import app


def create_pdf(page_count=3, label="Test") -> bytes:
    out = io.BytesIO()
    pdf = canvas.Canvas(out, pagesize=(300, 400))
    for index in range(page_count):
        pdf.setFont("Helvetica", 16)
        pdf.drawString(30, 350, f"{label} page {index + 1}")
        pdf.drawString(30, 320, "A small generated test document")
        pdf.showPage()
    pdf.save()
    return out.getvalue()


def create_table_pdf() -> bytes:
    out = io.BytesIO()
    pdf = canvas.Canvas(out, pagesize=(300, 400))
    xs = [30, 150, 270]
    ys = [260, 230, 200, 170]
    for x in xs:
        pdf.line(x, ys[0], x, ys[-1])
    for y in ys:
        pdf.line(xs[0], y, xs[-1], y)
    rows = [["Name", "Count"], ["Alpha", "10"], ["Beta", "20"]]
    for ri, row in enumerate(rows):
        for ci, value in enumerate(row):
            pdf.drawString(xs[ci] + 4, ys[ri + 1] + 8, value)
    pdf.save()
    return out.getvalue()


def create_text_pdf(lines, metadata=None) -> bytes:
    doc = fitz.open()
    page = doc.new_page(width=300, height=400)
    for index, line in enumerate(lines):
        page.insert_text((25, 50 + index * 32), line, fontsize=14)
    if metadata:
        doc.set_metadata(metadata)
    output = doc.tobytes()
    doc.close()
    return output


def create_form_pdf() -> bytes:
    doc = fitz.open()
    page = doc.new_page(width=300, height=400)
    configs = [
        ("full_name", fitz.PDF_WIDGET_TYPE_TEXT, fitz.Rect(30, 30, 200, 60)),
        ("agree", fitz.PDF_WIDGET_TYPE_CHECKBOX, fitz.Rect(30, 80, 50, 100)),
        ("country", fitz.PDF_WIDGET_TYPE_COMBOBOX, fitz.Rect(30, 120, 200, 150)),
    ]
    for name, field_type, rect in configs:
        widget = fitz.Widget()
        widget.field_name = name
        widget.field_type = field_type
        widget.rect = rect
        if name == "country":
            widget.choice_values = ["India", "Canada", "Other"]
            widget.field_value = "India"
        page.add_widget(widget)
    output = doc.tobytes()
    doc.close()
    return output


def download_bytes(client, url):
    response = client.get(url)
    if response.status_code != 200:
        raise AssertionError(f"Download failed: {response.status_code} {response.text}")
    return response.content


class PaperNestSmokeTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.client_context = TestClient(app)
        cls.client = cls.client_context.__enter__()

    @classmethod
    def tearDownClass(cls):
        cls.client_context.__exit__(None, None, None)

    def test_health(self):
        response = self.client.get("/api/health")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["status"], "ok")
        self.assertEqual(response.headers.get("cache-control"), "no-store")
        self.assertEqual(response.headers.get("x-content-type-options"), "nosniff")

    def test_merge_and_page_integrity(self):
        response = self.client.post("/api/process/merge-pdf", files=[
            ("files", ("a.pdf", create_pdf(2, "A"), "application/pdf")),
            ("files", ("b.pdf", create_pdf(1, "B"), "application/pdf")),
        ], data={"options": "{}"})
        self.assertEqual(response.status_code, 200, response.text)
        body = response.json()
        payload = download_bytes(self.client, body["download_url"])
        self.assertEqual(len(PdfReader(io.BytesIO(payload)).pages), 3)

    def test_remove_pages(self):
        response = self.client.post("/api/process/remove-pages", files=[("files", ("three.pdf", create_pdf(3), "application/pdf"))], data={"options": '{"pages":"2"}'})
        self.assertEqual(response.status_code, 200, response.text)
        data = download_bytes(self.client, response.json()["download_url"])
        self.assertEqual(len(PdfReader(io.BytesIO(data)).pages), 2)

    def test_extract_pages(self):
        response = self.client.post("/api/process/extract-pages", files=[("files", ("three.pdf", create_pdf(3), "application/pdf"))], data={"options": '{"pages":"1,3"}'})
        self.assertEqual(response.status_code, 200, response.text)
        data = download_bytes(self.client, response.json()["download_url"])
        self.assertEqual(len(PdfReader(io.BytesIO(data)).pages), 2)

    def test_rotate_pdf(self):
        response = self.client.post("/api/process/rotate-pdf", files=[("files", ("three.pdf", create_pdf(3), "application/pdf"))], data={"options": '{"pages":"2","rotation":90}'})
        self.assertEqual(response.status_code, 200, response.text)
        data = download_bytes(self.client, response.json()["download_url"])
        reader = PdfReader(io.BytesIO(data))
        self.assertEqual(len(reader.pages), 3)
        self.assertEqual(reader.pages[1].rotation, 90)

    def test_jpg_to_pdf(self):
        img_data = io.BytesIO()
        Image.new("RGB", (80, 120), "white").save(img_data, format="JPEG")
        response = self.client.post("/api/process/jpg-to-pdf", files=[("files", ("one.jpg", img_data.getvalue(), "image/jpeg"))], data={"options": '{"page_size":"a4"}'})
        self.assertEqual(response.status_code, 200, response.text)
        data = download_bytes(self.client, response.json()["download_url"])
        self.assertGreater(len(PdfReader(io.BytesIO(data)).pages), 0)

    def test_split_returns_zip(self):
        response = self.client.post("/api/process/split-pdf", files=[("files", ("three.pdf", create_pdf(3), "application/pdf"))], data={"options": '{"mode":"each"}'})
        self.assertEqual(response.status_code, 200, response.text)
        payload = download_bytes(self.client, response.json()["download_url"])
        with zipfile.ZipFile(io.BytesIO(payload)) as zf:
            self.assertEqual(len([name for name in zf.namelist() if name.endswith(".pdf")]), 3)
            for name in zf.namelist():
                self.assertEqual(len(PdfReader(io.BytesIO(zf.read(name))).pages), 1)

    def test_split_by_multiple_ranges(self):
        response = self.client.post("/api/process/split-pdf", files=[("files", ("six.pdf", create_pdf(6), "application/pdf"))], data={"options": '{"mode":"ranges","pages":"1-2,4-6"}'})
        self.assertEqual(response.status_code, 200, response.text)
        payload = download_bytes(self.client, response.json()["download_url"])
        with zipfile.ZipFile(io.BytesIO(payload)) as zf:
            outputs = [name for name in zf.namelist() if name.endswith(".pdf")]
            self.assertEqual(len(outputs), 2)
            self.assertEqual(sorted(len(PdfReader(io.BytesIO(zf.read(name))).pages) for name in outputs), [2, 3])

    def test_preview_has_page_count(self):
        response = self.client.post("/api/preview", files={"file": ("three.pdf", create_pdf(3), "application/pdf")})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["page_count"], 3)
        self.assertEqual(len(response.json()["page_dimensions"]), 3)
        self.assertGreaterEqual(len(response.json()["thumbnails"]), 1)


    def test_compress_pdf(self):
        response = self.client.post("/api/process/compress-pdf", files=[("files", ("three.pdf", create_pdf(3), "application/pdf"))], data={"options": '{"level":"recommended"}'})
        self.assertEqual(response.status_code, 200, response.text)
        data = download_bytes(self.client, response.json()["download_url"])
        self.assertEqual(len(PdfReader(io.BytesIO(data)).pages), 3)

    def test_pdf_to_jpg_zip(self):
        response = self.client.post("/api/process/pdf-to-jpg", files=[("files", ("three.pdf", create_pdf(3), "application/pdf"))], data={"options": '{"dpi":72}'})
        self.assertEqual(response.status_code, 200, response.text)
        payload = download_bytes(self.client, response.json()["download_url"])
        with zipfile.ZipFile(io.BytesIO(payload)) as zf:
            self.assertEqual(len(zf.namelist()), 3)
            self.assertTrue(all(name.endswith(".jpg") for name in zf.namelist()))

    def test_pdf_to_word_docx(self):
        response = self.client.post("/api/process/pdf-to-word", files=[("files", ("three.pdf", create_pdf(2), "application/pdf"))], data={"options": "{}"})
        self.assertEqual(response.status_code, 200, response.text)
        payload = download_bytes(self.client, response.json()["download_url"])
        with zipfile.ZipFile(io.BytesIO(payload)) as zf:
            self.assertIn("word/document.xml", zf.namelist())

    def test_pdf_to_excel_tables(self):
        response = self.client.post("/api/process/pdf-to-excel", files=[("files", ("table.pdf", create_table_pdf(), "application/pdf"))], data={"options": "{}"})
        self.assertEqual(response.status_code, 200, response.text)
        payload = download_bytes(self.client, response.json()["download_url"])
        workbook = load_workbook(io.BytesIO(payload))
        self.assertGreaterEqual(len(workbook.sheetnames), 1)

    def test_pdf_to_powerpoint_image_slides(self):
        response = self.client.post("/api/process/pdf-to-powerpoint", files=[("files", ("three.pdf", create_pdf(2), "application/pdf"))], data={"options": "{}"})
        self.assertEqual(response.status_code, 200, response.text)
        payload = download_bytes(self.client, response.json()["download_url"])
        presentation = Presentation(io.BytesIO(payload))
        self.assertEqual(len(presentation.slides), 2)

    def test_png_to_pdf(self):
        image_data = io.BytesIO()
        Image.new("RGBA", (80, 120), (50, 120, 240, 255)).save(image_data, format="PNG")
        response = self.client.post("/api/process/png-to-pdf", files=[("files", ("image.png", image_data.getvalue(), "image/png"))], data={"options": '{"page_size":"a4"}'})
        self.assertEqual(response.status_code, 200, response.text)
        payload = download_bytes(self.client, response.json()["download_url"])
        self.assertEqual(len(PdfReader(io.BytesIO(payload)).pages), 1)

    def test_editing_overlays_and_crop(self):
        cases = [
            ("page-numbers", '{"format":"Page {n} of {total}"}'),
            ("watermark-pdf", '{"text":"CONFIDENTIAL","font_size":18}'),
            ("add-text", '{"text":"Review copy","x":20,"y":25}'),
            ("edit-pdf", '{"text":"Note","shape":"rectangle","annotation":"note","annotation_text":"Review note","x":20,"y":20,"width":20,"height":10}'),
            ("crop-pdf", '{"top":4,"right":4,"bottom":4,"left":4}'),
        ]
        for tool, options in cases:
            with self.subTest(tool=tool):
                response = self.client.post(f"/api/process/{tool}", files=[("files", ("three.pdf", create_pdf(2), "application/pdf"))], data={"options": options})
                self.assertEqual(response.status_code, 200, response.text)
                payload = download_bytes(self.client, response.json()["download_url"])
                self.assertEqual(len(PdfReader(io.BytesIO(payload)).pages), 2)

    def test_organize_reorder_delete_rotate(self):
        response = self.client.post("/api/process/organize-pdf", files=[("files", ("three.pdf", create_pdf(3), "application/pdf"))], data={"options": '{"order":"3,1,2","delete_pages":"2","rotate_pages":"3","rotation":90}'})
        self.assertEqual(response.status_code, 200, response.text)
        payload = download_bytes(self.client, response.json()["download_url"])
        reader = PdfReader(io.BytesIO(payload))
        self.assertEqual(len(reader.pages), 2)
        self.assertEqual(reader.pages[0].rotation, 90)
        self.assertIn("page 3", reader.pages[0].extract_text().lower())
        self.assertIn("page 1", reader.pages[1].extract_text().lower())

    def test_edit_pdf_annotations(self):
        for annotation in ("note", "highlight", "underline", "strikeout"):
            with self.subTest(annotation=annotation):
                options = '{"text":"Marked text","shape":"none","annotation":"%s","annotation_text":"Review note","x":15,"y":15,"width":35,"height":12}' % annotation
                response = self.client.post("/api/process/edit-pdf", files=[("files", ("one.pdf", create_pdf(1), "application/pdf"))], data={"options": options})
                self.assertEqual(response.status_code, 200, response.text)
                payload = download_bytes(self.client, response.json()["download_url"])
                with fitz.open(stream=payload, filetype="pdf") as doc:
                    annotations = list(doc[0].annots() or [])
                    self.assertGreaterEqual(len(annotations), 1)

    def test_compare_and_repair(self):
        compare = self.client.post("/api/process/compare-pdf", files=[
            ("files", ("left.pdf", create_pdf(2, "Left"), "application/pdf")),
            ("files", ("right.pdf", create_pdf(3, "Right"), "application/pdf")),
        ], data={"options": '{"mode":"visual"}'})
        self.assertEqual(compare.status_code, 200, compare.text)
        comparison = download_bytes(self.client, compare.json()["download_url"])
        self.assertEqual(len(PdfReader(io.BytesIO(comparison)).pages), 3)
        repair = self.client.post("/api/process/repair-pdf", files=[("files", ("left.pdf", create_pdf(2), "application/pdf"))], data={"options": "{}"})
        self.assertEqual(repair.status_code, 200, repair.text)
        repaired = download_bytes(self.client, repair.json()["download_url"])
        self.assertEqual(len(PdfReader(io.BytesIO(repaired)).pages), 2)

    @unittest.skipUnless(shutil.which("libreoffice") or shutil.which("soffice"), "LibreOffice is required for office-to-PDF conversion")
    def test_office_to_pdf_conversions(self):
        doc = Document(); doc.add_paragraph("Word conversion smoke test")
        word = io.BytesIO(); doc.save(word)
        workbook = Workbook(); workbook.active["A1"] = "Excel smoke test"
        xlsx = io.BytesIO(); workbook.save(xlsx)
        presentation = Presentation(); slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "PowerPoint smoke test"
        pptx = io.BytesIO(); presentation.save(pptx)
        sources = [
            ("word-to-pdf", "sample.docx", word.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            ("excel-to-pdf", "sample.xlsx", xlsx.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
            ("powerpoint-to-pdf", "sample.pptx", pptx.getvalue(), "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
        ]
        for tool, name, data, mime in sources:
            with self.subTest(tool=tool):
                response = self.client.post(f"/api/process/{tool}", files=[("files", (name, data, mime))], data={"options": "{}"})
                self.assertEqual(response.status_code, 200, response.text)
                payload = download_bytes(self.client, response.json()["download_url"])
                self.assertGreaterEqual(len(PdfReader(io.BytesIO(payload)).pages), 1)

    def test_protect_and_unlock_pdf(self):
        protect = self.client.post("/api/process/protect-pdf", files=[("files", ("plain.pdf", create_pdf(2), "application/pdf"))], data={"options": '{"password":"test-password-123","password_confirm":"test-password-123"}'})
        self.assertEqual(protect.status_code, 200, protect.text)
        protected_bytes = download_bytes(self.client, protect.json()["download_url"])
        protected_reader = PdfReader(io.BytesIO(protected_bytes))
        self.assertTrue(protected_reader.is_encrypted)
        self.assertEqual(protected_reader.decrypt("test-password-123"), 2)
        self.assertEqual(len(protected_reader.pages), 2)

        unlock = self.client.post("/api/process/unlock-pdf", files=[("files", ("locked.pdf", protected_bytes, "application/pdf"))], data={"options": '{"password":"test-password-123"}'})
        self.assertEqual(unlock.status_code, 200, unlock.text)
        unlocked_bytes = download_bytes(self.client, unlock.json()["download_url"])
        unlocked_reader = PdfReader(io.BytesIO(unlocked_bytes))
        self.assertFalse(unlocked_reader.is_encrypted)
        self.assertEqual(len(unlocked_reader.pages), 2)

    def test_unlock_wrong_password_is_rejected(self):
        writer = PdfReader(io.BytesIO(create_pdf(1)))
        from pypdf import PdfWriter
        encrypted = PdfWriter()
        for page in writer.pages: encrypted.add_page(page)
        encrypted.encrypt("correct-pass-123", algorithm="AES-256")
        out = io.BytesIO(); encrypted.write(out)
        response = self.client.post("/api/process/unlock-pdf", files=[("files", ("locked.pdf", out.getvalue(), "application/pdf"))], data={"options": '{"password":"incorrect-pass"}'})
        self.assertEqual(response.status_code, 400)

    def test_pdf_to_png_zip(self):
        response = self.client.post("/api/process/pdf-to-png", files=[("files", ("three.pdf", create_pdf(2), "application/pdf"))], data={"options": '{"dpi":72}'})
        self.assertEqual(response.status_code, 200, response.text)
        payload = download_bytes(self.client, response.json()["download_url"])
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            self.assertEqual(len(archive.namelist()), 2)
            self.assertTrue(all(name.endswith(".png") for name in archive.namelist()))

    @unittest.skipUnless(shutil.which("tesseract"), "Tesseract OCR is required for OCR conversion")
    def test_ocr_pdf_searchable_output(self):
        font_path = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
        from PIL import ImageDraw, ImageFont
        image = Image.new("RGB", (1200, 350), "white")
        draw = ImageDraw.Draw(image)
        font = ImageFont.truetype(font_path, 58) if Path(font_path).exists() else ImageFont.load_default()
        draw.text((65, 100), "PAPERNEST OCR TEST", fill="black", font=font)
        image_buf = io.BytesIO(); image.save(image_buf, format="PNG")
        scanned = fitz.open()
        page = scanned.new_page(width=432, height=126)
        page.insert_image(page.rect, stream=image_buf.getvalue())
        pdf_buf = io.BytesIO(); scanned.save(pdf_buf); scanned.close()
        response = self.client.post("/api/process/ocr-pdf", files=[("files", ("scan.pdf", pdf_buf.getvalue(), "application/pdf"))], data={"options": '{"language":"eng"}'})
        self.assertEqual(response.status_code, 200, response.text)
        output = download_bytes(self.client, response.json()["download_url"])
        with fitz.open(stream=output, filetype="pdf") as document:
            extracted = document[0].get_text().upper()
            self.assertIn("PAPERNEST", extracted)
            self.assertIn("OCR", extracted)

    def test_pdf_to_text_rejects_scan_without_ocr(self):
        image_buf = io.BytesIO()
        Image.new("RGB", (120, 120), "white").save(image_buf, format="PNG")
        scanned = fitz.open(); page = scanned.new_page(width=120, height=120); page.insert_image(page.rect, stream=image_buf.getvalue())
        pdf_buf = io.BytesIO(); scanned.save(pdf_buf); scanned.close()
        response = self.client.post("/api/process/pdf-to-text", files=[("files", ("scan.pdf", pdf_buf.getvalue(), "application/pdf"))], data={"options": "{}"})
        self.assertEqual(response.status_code, 422)
        self.assertIn("OCR", response.json()["detail"])

    def test_pdf_to_text(self):
        response = self.client.post("/api/process/pdf-to-text", files=[("files", ("text.pdf", create_pdf(2, "Extract me"), "application/pdf"))], data={"options": "{}"})
        self.assertEqual(response.status_code, 200, response.text)
        payload = download_bytes(self.client, response.json()["download_url"])
        text = payload.decode("utf-8")
        self.assertIn("Extract me page 1", text)
        self.assertIn("===== Page 2 =====", text)

    def test_pdf_metadata_update(self):
        response = self.client.post("/api/process/pdf-metadata", files=[("files", ("source.pdf", create_pdf(1), "application/pdf"))], data={"options": '{"title":"Quarterly Report","author":"PaperNest QA"}'})
        self.assertEqual(response.status_code, 200, response.text)
        payload = download_bytes(self.client, response.json()["download_url"])
        with fitz.open(stream=payload, filetype="pdf") as document:
            self.assertEqual(document.metadata.get("title"), "Quarterly Report")
            self.assertEqual(document.metadata.get("author"), "PaperNest QA")

    def test_sign_pdf_image_overlay(self):
        image_data = io.BytesIO()
        Image.new("RGBA", (120, 40), (10, 20, 80, 180)).save(image_data, format="PNG")
        response = self.client.post("/api/process/sign-pdf", files=[
            ("files", ("source.pdf", create_pdf(2), "application/pdf")),
            ("files", ("signature.png", image_data.getvalue(), "image/png")),
        ], data={"options": '{"pages":"2","x":40,"y":70,"width":25,"height":10}'})
        self.assertEqual(response.status_code, 200, response.text)
        payload = download_bytes(self.client, response.json()["download_url"])
        with fitz.open(stream=payload, filetype="pdf") as document:
            self.assertEqual(document.page_count, 2)
            self.assertEqual(len(document[1].get_images()), 1)

    def test_redact_visual_preview_session_and_page(self):
        response = self.client.post("/api/redact-preview", files={"file": ("three.pdf", create_pdf(3), "application/pdf")})
        self.assertEqual(response.status_code, 200, response.text)
        session = response.json()
        self.assertEqual(session["page_count"], 3)
        session_id = session["session_id"]
        self.assertRegex(session_id, r"^[a-f0-9]{32}$")
        page_response = self.client.get(f"/api/redact-preview/{session_id}/page/2")
        self.assertEqual(page_response.status_code, 200, page_response.text)
        page = page_response.json()
        self.assertEqual(page["page"], 2)
        self.assertGreater(page["width_pt"], 0)
        self.assertGreater(page["height_pt"], 0)
        self.assertTrue(page["data_url"].startswith("data:image/png;base64,"))
        self.assertEqual(self.client.get(f"/api/redact-preview/{session_id}/page/99").status_code, 404)

    def test_redact_preview_rejects_non_pdf(self):
        response = self.client.post("/api/redact-preview", files={"file": ("bad.txt", b"not a pdf", "text/plain")})
        self.assertEqual(response.status_code, 400)

    def test_redact_pdf_removes_text_and_metadata(self):
        source = create_text_pdf(["SECRET-123", "Keep this line"], {"title": "Sensitive title", "author": "Private author"})
        response = self.client.post("/api/process/redact-pdf", files=[("files", ("private.pdf", source, "application/pdf"))], data={"options": '{"terms":"SECRET-123","confirmed":"yes"}'})
        self.assertEqual(response.status_code, 200, response.text)
        output = download_bytes(self.client, response.json()["download_url"])
        with fitz.open(stream=output, filetype="pdf") as document:
            text = document[0].get_text()
            self.assertNotIn("SECRET-123", text)
            self.assertIn("Keep this line", text)
            self.assertEqual(document.metadata.get("title"), "")
            self.assertEqual(document.metadata.get("author"), "")
            pix = document[0].get_pixmap(matrix=fitz.Matrix(1, 1), alpha=False)
            self.assertGreater(len(pix.samples), 0)

    def test_redact_pdf_visual_region_removes_underlying_text(self):
        # Coordinates are represented in the visual page coordinate system used by the mouse selector.
        options = {"confirmed": "yes", "boxes_json": json.dumps([{"page": 1, "x0": 15, "y0": 18, "x1": 220, "y1": 80, "visual": True}]), "terms": "", "fill": "black"}
        response = self.client.post("/api/process/redact-pdf", files=[("files", ("three.pdf", create_pdf(1, "SecretName"), "application/pdf"))], data={"options": json.dumps(options)})
        self.assertEqual(response.status_code, 200, response.text)
        payload = download_bytes(self.client, response.json()["download_url"])
        with fitz.open(stream=payload, filetype="pdf") as doc:
            self.assertNotIn("SecretName page 1", doc[0].get_text())

    def test_redact_pdf_blur_removes_text_and_renders_blurred_region(self):
        source = create_text_pdf(["BLUR-ME-SECRET", "Keep this line"])
        options = {"confirmed": "yes", "boxes_json": json.dumps([{"page": 1, "x0": 15, "y0": 30, "x1": 180, "y1": 65}]), "fill": "blur"}
        response = self.client.post("/api/process/redact-pdf", files=[("files", ("blur.pdf", source, "application/pdf"))], data={"options": json.dumps(options)})
        self.assertEqual(response.status_code, 200, response.text)
        output = download_bytes(self.client, response.json()["download_url"])
        with fitz.open(stream=output, filetype="pdf") as document:
            self.assertNotIn("BLUR-ME-SECRET", document[0].get_text())
            self.assertIn("Keep this line", document[0].get_text())
            self.assertGreater(len(document[0].get_images()), 0)

    def test_redact_pdf_rejects_unknown_fill_style(self):
        source = create_text_pdf(["Private line"])
        options = {"confirmed": "yes", "boxes_json": json.dumps([{"page": 1, "x0": 15, "y0": 30, "x1": 180, "y1": 65}]), "fill": "rainbow"}
        response = self.client.post("/api/process/redact-pdf", files=[("files", ("private.pdf", source, "application/pdf"))], data={"options": json.dumps(options)})
        self.assertEqual(response.status_code, 400)

    def test_redact_pdf_visual_region_handles_rotated_page(self):
        source = create_pdf(1, "RotatedSecret")
        with fitz.open(stream=source, filetype="pdf") as document:
            page = document[0]
            original_box = page.search_for("RotatedSecret page 1")[0]
            page.set_rotation(90)
            visible_box = original_box * page.rotation_matrix
            visible_box = fitz.Rect(max(0, visible_box.x0 - 2), max(0, visible_box.y0 - 2), min(page.rect.width, visible_box.x1 + 2), min(page.rect.height, visible_box.y1 + 2))
            rotated_source = document.tobytes()
        options = {"confirmed": "yes", "boxes_json": json.dumps([{"page": 1, "x0": visible_box.x0, "y0": visible_box.y0, "x1": visible_box.x1, "y1": visible_box.y1, "visual": True}]), "terms": "", "fill": "black"}
        response = self.client.post("/api/process/redact-pdf", files=[("files", ("rotated.pdf", rotated_source, "application/pdf"))], data={"options": json.dumps(options)})
        self.assertEqual(response.status_code, 200, response.text)
        payload = download_bytes(self.client, response.json()["download_url"])
        with fitz.open(stream=payload, filetype="pdf") as document:
            self.assertNotIn("RotatedSecret page 1", document[0].get_text())
            self.assertEqual(document[0].rotation, 90)

    def test_redact_pdf_rectangle_removes_underlying_text(self):
        source = create_text_pdf(["CONFIDENTIAL REGION", "Keep this line"])
        boxes = [{"page": 1, "x0": 15, "y0": 30, "x1": 180, "y1": 65}]
        response = self.client.post("/api/process/redact-pdf", files=[("files", ("private.pdf", source, "application/pdf"))], data={"options": json.dumps({"boxes_json": json.dumps(boxes), "confirmed": "yes"})})
        self.assertEqual(response.status_code, 200, response.text)
        output = download_bytes(self.client, response.json()["download_url"])
        with fitz.open(stream=output, filetype="pdf") as document:
            self.assertNotIn("CONFIDENTIAL REGION", document[0].get_text())
            self.assertIn("Keep this line", document[0].get_text())

    def test_redaction_requires_confirmation_and_match(self):
        source = create_text_pdf(["Visible content"])
        not_confirmed = self.client.post("/api/process/redact-pdf", files=[("files", ("a.pdf", source, "application/pdf"))], data={"options": '{"terms":"Visible content"}'})
        self.assertEqual(not_confirmed.status_code, 400)
        no_match = self.client.post("/api/process/redact-pdf", files=[("files", ("a.pdf", source, "application/pdf"))], data={"options": '{"terms":"NOT-FOUND","confirmed":"yes"}'})
        self.assertEqual(no_match.status_code, 422)

    def test_form_fields_detect_and_fill(self):
        source = create_form_pdf()
        inspect = self.client.post("/api/form-fields", files={"file": ("form.pdf", source, "application/pdf")})
        self.assertEqual(inspect.status_code, 200, inspect.text)
        self.assertEqual({item["name"] for item in inspect.json()["fields"]}, {"full_name", "agree", "country"})
        response = self.client.post("/api/process/fill-pdf-form", files=[("files", ("form.pdf", source, "application/pdf"))], data={"options": json.dumps({"form_values": {"full_name": "Alice Example", "agree": True, "country": "Canada"}, "flatten": "no"})})
        self.assertEqual(response.status_code, 200, response.text)
        output = download_bytes(self.client, response.json()["download_url"])
        with fitz.open(stream=output, filetype="pdf") as document:
            fields = {widget.field_name: widget.field_value for page in document for widget in (page.widgets() or [])}
        self.assertEqual(fields["full_name"], "Alice Example")
        self.assertEqual(fields["agree"], "Yes")
        self.assertEqual(fields["country"], "Canada")

    def test_form_fill_can_flatten_fields(self):
        source = create_form_pdf()
        response = self.client.post("/api/process/fill-pdf-form", files=[("files", ("form.pdf", source, "application/pdf"))], data={"options": json.dumps({"form_values": {"full_name": "Flattened Name"}, "flatten": "yes"})})
        self.assertEqual(response.status_code, 200, response.text)
        output = download_bytes(self.client, response.json()["download_url"])
        with fitz.open(stream=output, filetype="pdf") as document:
            self.assertEqual(sum(1 for page in document for _ in (page.widgets() or [])), 0)
            self.assertIn("Flattened Name", document[0].get_text())

    def test_compare_pdf_text_diffs_and_combined_mode(self):
        left = create_text_pdf(["Title", "Old value", "Same line"])
        right = create_text_pdf(["Title", "New value", "Same line", "New extra line"])
        response = self.client.post("/api/process/compare-pdf", files=[("files", ("left.pdf", left, "application/pdf")), ("files", ("right.pdf", right, "application/pdf"))], data={"options": '{"mode":"both"}'})
        self.assertEqual(response.status_code, 200, response.text)
        output = download_bytes(self.client, response.json()["download_url"])
        with fitz.open(stream=output, filetype="pdf") as document:
            text = "\n".join(page.get_text() for page in document)
            self.assertGreaterEqual(document.page_count, 3)
            self.assertIn("Old value", text)
            self.assertIn("New value", text)
            self.assertIn("New extra line", text)

    def test_organize_preview_returns_every_page_thumbnail(self):
        response = self.client.post("/api/organize-preview", files={"file": ("six.pdf", create_pdf(6), "application/pdf")})
        self.assertEqual(response.status_code, 200, response.text)
        payload = response.json()
        self.assertEqual(payload["page_count"], 6)
        self.assertEqual([page["page"] for page in payload["pages"]], list(range(1, 7)))
        self.assertTrue(all(page["data_url"].startswith("data:image/jpeg;base64,") for page in payload["pages"]))
        self.assertTrue(all(page["width_pt"] > 0 and page["height_pt"] > 0 for page in payload["pages"]))

    def test_html_to_pdf_url_only_request(self):
        def fake_url_renderer(paths, options, outdir, base):
            self.assertEqual(paths, [])
            self.assertEqual(options.get("url"), "https://example.com/report")
            outdir.mkdir(parents=True, exist_ok=True)
            target = outdir / f"{base}.pdf"
            target.write_bytes(create_pdf(1, "Public URL page"))
            return target, "URL converted in test renderer."

        with patch("app.main._html_to_pdf", side_effect=fake_url_renderer):
            response = self.client.post("/api/process/html-to-pdf", data={"options": json.dumps({"url": "https://example.com/report", "page_format": "A4"})})
        self.assertEqual(response.status_code, 200, response.text)
        payload = download_bytes(self.client, response.json()["download_url"])
        with fitz.open(stream=payload, filetype="pdf") as document:
            self.assertEqual(document.page_count, 1)
            self.assertIn("Public URL page", document[0].get_text())

    def test_html_to_pdf_uploaded_html(self):
        html = b"<!doctype html><html><head><title>Smoke test</title><style>h1{color:#c65f43}</style></head><body><h1>PaperNest HTML TEST</h1><p>HTML gets rendered to a real PDF.</p></body></html>"
        response = self.client.post("/api/process/html-to-pdf", files=[("files", ("sample.html", html, "text/html"))], data={"options": '{"page_format":"A4","margin_mm":10,"print_background":"yes"}'})
        self.assertEqual(response.status_code, 200, response.text)
        payload = download_bytes(self.client, response.json()["download_url"])
        with fitz.open(stream=payload, filetype="pdf") as document:
            self.assertEqual(document.page_count, 1)
            self.assertIn("PaperNest HTML TEST", document[0].get_text())

    def test_html_to_pdf_rejects_uploaded_html_and_url_together(self):
        response = self.client.post("/api/process/html-to-pdf", files=[("files", ("page.html", b"<html><body>Hi</body></html>", "text/html"))], data={"options": '{"url":"https://example.com"}'})
        self.assertEqual(response.status_code, 400, response.text)
        self.assertIn("either a webpage URL or an uploaded HTML file", response.text)

    def test_html_to_pdf_rejects_private_or_non_http_urls(self):
        for url in ("http://127.0.0.1/", "http://localhost/", "file:///etc/passwd"):
            with self.subTest(url=url):
                response = self.client.post("/api/process/html-to-pdf", data={"options": json.dumps({"url": url})})
                self.assertEqual(response.status_code, 400, response.text)

    def test_sanitize_pdf_metadata_only(self):
        source = create_text_pdf(["Visible content"], {"title": "Secret title", "author": "Secret author", "subject": "Private subject"})
        response = self.client.post("/api/process/sanitize-pdf-metadata", files=[("files", ("secret.pdf", source, "application/pdf"))], data={"options": '{"scope":"metadata"}'})
        self.assertEqual(response.status_code, 200, response.text)
        payload = download_bytes(self.client, response.json()["download_url"])
        with fitz.open(stream=payload, filetype="pdf") as document:
            self.assertEqual(document.page_count, 1)
            self.assertIn("Visible content", document[0].get_text())
            self.assertEqual(document.metadata.get("title"), "")
            self.assertEqual(document.metadata.get("author"), "")
            self.assertEqual(document.metadata.get("subject"), "")

    def test_sanitize_pdf_full_scope_removes_annotations_and_metadata(self):
        source = create_text_pdf(["Visible content"], {"title": "Secret title"})
        with fitz.open(stream=source, filetype="pdf") as document:
            document[0].add_text_annot((30, 70), "Confidential review comment")
            source = document.tobytes()
        response = self.client.post("/api/process/sanitize-pdf-metadata", files=[("files", ("secret.pdf", source, "application/pdf"))], data={"options": '{"scope":"full"}'})
        self.assertEqual(response.status_code, 200, response.text)
        payload = download_bytes(self.client, response.json()["download_url"])
        with fitz.open(stream=payload, filetype="pdf") as document:
            self.assertIn("Visible content", document[0].get_text())
            self.assertEqual(document.metadata.get("title"), "")
            self.assertEqual(list(document[0].annots() or []), [])

    def test_rejects_mislabeled_pdf(self):
        response = self.client.post("/api/process/rotate-pdf", files=[("files", ("bad.pdf", b"not a pdf", "application/pdf"))], data={"options": "{}"})
        self.assertEqual(response.status_code, 400)


if __name__ == "__main__":
    unittest.main()
