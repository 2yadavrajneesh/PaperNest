from __future__ import annotations

import asyncio
import base64
import difflib
import io
import ipaddress
import json
import os
import re
import shutil
import socket
import subprocess
import tempfile
import time
import uuid
import zipfile
from contextlib import asynccontextmanager
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit

import fitz
import pdfplumber
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from starlette.concurrency import run_in_threadpool
from openpyxl import Workbook
from PIL import Image, ImageFilter
from pypdf import PdfReader, PdfWriter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

APP_NAME = "PaperNest API"
ROOT = Path(os.getenv("PAPERNEST_TEMP_DIR", "/tmp/papernest-jobs")).resolve()
ROOT.mkdir(parents=True, exist_ok=True, mode=0o700)
try:
    ROOT.chmod(0o700)
except OSError:
    pass
MAX_FILE_BYTES = int(os.getenv("MAX_FILE_BYTES", str(50 * 1024 * 1024)))
MAX_TOTAL_BYTES = int(os.getenv("MAX_TOTAL_BYTES", str(100 * 1024 * 1024)))
MAX_OUTPUT_BYTES = int(os.getenv("MAX_OUTPUT_BYTES", str(250 * 1024 * 1024)))
MAX_FILES = int(os.getenv("MAX_FILES", "20"))
RETENTION_SECONDS = int(os.getenv("RETENTION_SECONDS", "1800"))
MAX_PDF_PAGES = int(os.getenv("MAX_PDF_PAGES", "500"))
MAX_OCR_PAGES = int(os.getenv("MAX_OCR_PAGES", "100"))
ALLOWED_ORIGINS = [x.strip() for x in os.getenv("ALLOWED_ORIGINS", "http://localhost:3000,http://127.0.0.1:3000").split(",") if x.strip()]

PDF_TOOLS = {
    "merge-pdf", "split-pdf", "compress-pdf", "organize-pdf", "rotate-pdf",
    "remove-pages", "extract-pages", "pdf-to-word", "word-to-pdf", "pdf-to-jpg",
    "jpg-to-pdf", "pdf-to-excel", "excel-to-pdf", "pdf-to-powerpoint",
    "powerpoint-to-pdf", "png-to-pdf", "edit-pdf", "page-numbers", "watermark-pdf",
    "crop-pdf", "add-text", "repair-pdf", "compare-pdf", "pdf-reader",
    "protect-pdf", "unlock-pdf", "pdf-to-png", "pdf-to-text", "pdf-metadata", "sign-pdf", "ocr-pdf",
    "redact-pdf", "fill-pdf-form", "html-to-pdf", "sanitize-pdf-metadata",
}
PDF_BATCHABLE = {
    "split-pdf", "compress-pdf", "organize-pdf", "rotate-pdf", "remove-pages",
    "extract-pages", "pdf-to-word", "pdf-to-jpg", "pdf-to-excel", "pdf-to-powerpoint",
    "edit-pdf", "page-numbers", "watermark-pdf", "crop-pdf", "add-text", "repair-pdf",
    "word-to-pdf", "excel-to-pdf", "powerpoint-to-pdf", "pdf-to-png", "pdf-to-text", "protect-pdf", "unlock-pdf", "ocr-pdf", "sanitize-pdf-metadata",
}
IMAGE_EXTS = {".jpg", ".jpeg", ".png"}
PDF_EXTS = {".pdf"}
WORD_EXTS = {".docx", ".doc"}
EXCEL_EXTS = {".xlsx", ".xls"}
PPT_EXTS = {".pptx", ".ppt"}
HTML_EXTS = {".html", ".htm"}
ALL_EXTS = PDF_EXTS | IMAGE_EXTS | WORD_EXTS | EXCEL_EXTS | PPT_EXTS | HTML_EXTS
MIME_TYPES = {
    ".pdf": "application/pdf", ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".doc": "application/msword", ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".xls": "application/vnd.ms-excel", ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".ppt": "application/vnd.ms-powerpoint", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
    ".png": "image/png", ".html": "text/html; charset=utf-8", ".htm": "text/html; charset=utf-8", ".txt": "text/plain; charset=utf-8", ".zip": "application/zip",
}


def safe_stem(name: str) -> str:
    stem = Path(Path(name).name).stem
    stem = re.sub(r"[^A-Za-z0-9._-]+", "_", stem).strip("._-")[:80]
    return stem or "document"


def _file_signature_ok(path: Path, data_start: bytes) -> bool:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return b"%PDF-" in data_start[:1024]
    if ext in {".jpg", ".jpeg"}:
        return data_start.startswith(b"\xff\xd8\xff")
    if ext == ".png":
        return data_start.startswith(b"\x89PNG\r\n\x1a\n")
    if ext in HTML_EXTS:
        sample = data_start.lstrip(b"\xef\xbb\xbf \t\r\n").lower()
        return b"<" in sample and b">" in sample and b"\x00" not in sample
    if ext in {".docx", ".xlsx", ".pptx"}:
        return data_start.startswith(b"PK\x03\x04") or data_start.startswith(b"PK\x05\x06")
    if ext in {".doc", ".xls", ".ppt"}:
        return data_start.startswith(bytes.fromhex("D0CF11E0A1B11AE1"))
    return False


def _check_file(path: Path, content_start: bytes) -> None:
    ext = path.suffix.lower()
    if ext not in ALL_EXTS:
        raise HTTPException(400, f"Unsupported file type: {path.suffix or 'unknown'}. Choose a supported document or image file.")
    if not _file_signature_ok(path, content_start):
        raise HTTPException(400, f"{path.name} does not look like a valid {ext} file. The file may be mislabeled or damaged.")
    # OOXML files are ZIP containers. Bound their expanded size and entry count to reduce zip-bomb risk.
    if ext in {".docx", ".xlsx", ".pptx"}:
        required = {
            ".docx": {"[Content_Types].xml", "word/document.xml"},
            ".xlsx": {"[Content_Types].xml", "xl/workbook.xml"},
            ".pptx": {"[Content_Types].xml", "ppt/presentation.xml"},
        }[ext]
        try:
            with zipfile.ZipFile(path) as archive:
                members = archive.infolist()
                if len(members) > 10000:
                    raise HTTPException(413, "Office document contains too many internal entries to process safely.")
                expanded = sum(member.file_size for member in members)
                if expanded > 250 * 1024 * 1024:
                    raise HTTPException(413, "Office document expands beyond the 250 MB safety limit.")
                if not required.issubset(set(archive.namelist())):
                    raise HTTPException(400, f"{path.name} is not a structurally valid {ext} document.")
        except HTTPException:
            raise
        except (zipfile.BadZipFile, OSError) as exc:
            raise HTTPException(400, f"{path.name} is a damaged Office document.") from exc


def _parse_options(raw: str) -> dict[str, Any]:
    try:
        value = json.loads(raw or "{}")
    except json.JSONDecodeError as exc:
        raise HTTPException(400, "Tool settings are invalid JSON.") from exc
    if not isinstance(value, dict):
        raise HTTPException(400, "Tool settings must be an object.")
    return value


def _read_pdf(path: Path) -> PdfReader:
    try:
        reader = PdfReader(str(path), strict=False)
        if reader.is_encrypted:
            raise HTTPException(422, f"{path.name} is password-protected. Unlock it before processing.")
        if len(reader.pages) < 1:
            raise HTTPException(422, f"{path.name} contains no pages.")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise HTTPException(413, f"{path.name} has more than the {MAX_PDF_PAGES}-page processing limit.")
        return reader
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, f"Could not read {path.name}. The PDF may be corrupted or unsupported.") from exc


def _parse_pages(spec: str | None, total: int, *, default_all: bool = False) -> list[int]:
    """Parse a 1-indexed page list such as '1-3,5' into zero-indexed positions."""
    if spec is None or not str(spec).strip():
        return list(range(total)) if default_all else []
    selected: list[int] = []
    try:
        for part in str(spec).split(","):
            part = part.strip()
            if not part:
                continue
            if "-" in part:
                left, right = part.split("-", 1)
                start, end = int(left.strip()), int(right.strip())
                if start < 1 or end < start or end > total:
                    raise ValueError
                selected.extend(range(start - 1, end))
            else:
                page = int(part)
                if page < 1 or page > total:
                    raise ValueError
                selected.append(page - 1)
    except ValueError as exc:
        raise HTTPException(400, f"Invalid page range '{spec}'. Use page numbers such as 1-3,5; this PDF has {total} pages.") from exc
    # Keep user order while removing duplicates.
    result = list(dict.fromkeys(selected))
    if not result:
        raise HTTPException(400, "Select at least one page to process.")
    return result


def _write_pdf(writer: PdfWriter, out: Path) -> Path:
    try:
        with out.open("wb") as f:
            writer.write(f)
    except Exception as exc:
        raise HTTPException(422, "The PDF engine could not write the output document.") from exc
    return out


def _make_zip(files: list[tuple[Path, str]], out: Path) -> Path:
    with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as zf:
        for path, arcname in files:
            zf.write(path, arcname=arcname)
    return out


def _run_libreoffice(source: Path, output_dir: Path) -> Path:
    # Use the same Windows-aware detector that reads LIBREOFFICE_PATH.
    # The previous implementation only used shutil.which(), which does not
    # find LibreOffice when it is installed but not added to Windows PATH.
    executable = _find_libreoffice()
    if not executable:
        configured = os.getenv("LIBREOFFICE_PATH", "").strip().strip('"')
        raise HTTPException(
            503,
            "LibreOffice was not detected by the PaperNest API. "
            "Set LIBREOFFICE_PATH to the full path of soffice.exe and restart the API. "
            f"Configured path: {configured or '(not set)'}"
        )
    output_dir.mkdir(parents=True, exist_ok=True)
    profile = output_dir / f"lo-profile-{uuid.uuid4().hex}"
    command = [executable, f"-env:UserInstallation={profile.as_uri()}", "--headless", "--convert-to", "pdf", "--outdir", str(output_dir), str(source)]
    try:
        proc = subprocess.run(command, capture_output=True, timeout=180, check=False)
    except subprocess.TimeoutExpired as exc:
        raise HTTPException(504, "Office conversion exceeded the 3-minute processing limit.") from exc
    candidate = output_dir / f"{source.stem}.pdf"
    if proc.returncode != 0 or not candidate.exists() or candidate.stat().st_size == 0:
        # Never include document text or password-bearing arguments in the error response.
        detail = proc.stderr.decode("utf-8", errors="ignore")[-300:]
        raise HTTPException(422, f"LibreOffice could not convert this file. Check that it is valid and not password-protected. {detail}")
    return candidate

def _find_libreoffice() -> str | None:
    """Find LibreOffice on Windows, including registry and custom install paths."""
    import os
    import shutil
    import subprocess
    from pathlib import Path

    candidates = []

    # 1. Explicit environment variable
    configured = os.getenv("LIBREOFFICE_PATH", "").strip().strip('"')
    if configured:
        candidates.append(Path(configured).expanduser())

    # 2. Search PATH
    for name in ("soffice.exe", "soffice", "libreoffice.exe", "libreoffice"):
        found = shutil.which(name)
        if found:
            candidates.append(Path(found))

    # 3. Standard Windows installation folders
    for variable in ("PROGRAMFILES", "PROGRAMFILES(X86)", "LOCALAPPDATA"):
        base = os.getenv(variable)
        if base:
            candidates.extend([
                Path(base) / "LibreOffice" / "program" / "soffice.exe",
                Path(base) / "Programs" / "LibreOffice" / "program" / "soffice.exe",
            ])

    # 4. Check Windows Registry for the actual installation directory
    if os.name == "nt":
        try:
            import winreg

            registry_locations = [
                (winreg.HKEY_LOCAL_MACHINE,
                 r"SOFTWARE\LibreOffice\UNO\InstallPath"),
                (winreg.HKEY_LOCAL_MACHINE,
                 r"SOFTWARE\WOW6432Node\LibreOffice\UNO\InstallPath"),
                (winreg.HKEY_CURRENT_USER,
                 r"SOFTWARE\LibreOffice\UNO\InstallPath"),
            ]

            for hive, key_name in registry_locations:
                try:
                    with winreg.OpenKey(hive, key_name) as key:
                        install_dir, _ = winreg.QueryValueEx(key, "")
                        candidates.append(
                            Path(install_dir) / "soffice.exe"
                        )
                except OSError:
                    pass
        except ImportError:
            pass

    # 5. Return the first existing executable
    seen = set()
    for candidate in candidates:
        try:
            candidate = candidate.resolve()
            if str(candidate).lower() in seen:
                continue
            seen.add(str(candidate).lower())

            if candidate.is_file():
                # Verify that Windows can actually launch it
                result = subprocess.run(
                    [str(candidate), "--headless", "--version"],
                    capture_output=True,
                    text=True,
                    timeout=15,
                    check=False,
                )
                if result.returncode == 0:
                    return str(candidate)
        except (OSError, subprocess.SubprocessError):
            continue

    return None

def _pdf_to_docx(source: Path, out: Path) -> str:
    try:
        from pdf2docx import Converter  # type: ignore
        converter = Converter(str(source))
        try:
            converter.convert(str(out))
        finally:
            converter.close()
        if out.exists() and out.stat().st_size:
            return "Converted with the layout-based PDF-to-Word engine. Review complex layouts before editing."
    except ImportError:
        pass
    except Exception:
        # Text fallback remains useful for PDFs the optional layout engine cannot parse.
        pass
    from docx import Document
    try:
        pdf = fitz.open(str(source))
        doc = Document()
        for index, page in enumerate(pdf):
            text = page.get_text("text", sort=True).strip()
            if text:
                for line in text.splitlines():
                    if line.strip():
                        doc.add_paragraph(line)
            else:
                doc.add_paragraph(f"[No selectable text found on PDF page {index + 1}. This page may be scanned.]", style="Caption")
            if index < len(pdf) - 1:
                doc.add_page_break()
        doc.save(str(out))
    except Exception as exc:
        raise HTTPException(422, "Could not extract text from this PDF. It may be scanned, encrypted, or damaged.") from exc
    return "Created an editable DOCX from extracted text. Complex layout, images, columns, and scanned pages may not be preserved; OCR is not automatic."


def _pdf_to_xlsx(source: Path, out: Path) -> str:
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    sheets = 0
    try:
        with pdfplumber.open(str(source)) as pdf:
            for page_index, page in enumerate(pdf.pages, start=1):
                tables = page.extract_tables()
                for table_index, table in enumerate(tables, start=1):
                    cleaned = [[("" if cell is None else str(cell)).strip() for cell in row] for row in table]
                    while cleaned and not any(cleaned[0]):
                        cleaned.pop(0)
                    if not cleaned:
                        continue
                    ws = wb.create_sheet(title=f"P{page_index}_T{table_index}"[:31])
                    for row in cleaned:
                        ws.append(row)
                    ws.freeze_panes = "A2"
                    sheets += 1
    except Exception as exc:
        raise HTTPException(422, "Could not analyze tables in this PDF. The file may be scanned or malformed.") from exc
    if not sheets:
        raise HTTPException(422, "No machine-detectable tables were found. Scanned pages and visually aligned text may need OCR or manual cleanup.")
    wb.save(out)
    return f"Extracted {sheets} table(s) to XLSX. Verify column alignment and merged cells in the result."


def _compress_pdf(source: Path, out: Path, level: str) -> str:
    presets = {"strong": "/screen", "recommended": "/ebook", "less": "/printer"}
    preset = presets.get(level, "/ebook")
    gs = shutil.which("gs")
    if gs:
        cmd = [gs, "-sDEVICE=pdfwrite", "-dCompatibilityLevel=1.5", f"-dPDFSETTINGS={preset}", "-dNOPAUSE", "-dQUIET", "-dBATCH", "-dSAFER", f"-sOutputFile={out}", str(source)]
        try:
            proc = subprocess.run(cmd, capture_output=True, timeout=180, check=False)
            if proc.returncode == 0 and out.exists() and out.stat().st_size:
                before, after = source.stat().st_size, out.stat().st_size
                return f"Compression complete ({before / 1024:.0f} KB → {after / 1024:.0f} KB). Results depend on the document's content; files with already-compressed images may not shrink."
        except subprocess.TimeoutExpired:
            pass
    try:
        doc = fitz.open(str(source))
        doc.save(str(out), garbage=4, deflate=True, clean=True)
        doc.close()
        return "Applied structural PDF optimization. Ghostscript was unavailable or could not optimize this file, so image recompression may be limited."
    except Exception as exc:
        raise HTTPException(422, "Could not optimize this PDF.") from exc


def _render_pdf_pages(source: Path, output_dir: Path, dpi: int = 144, ext: str = "jpg") -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    try:
        doc = fitz.open(str(source))
        if doc.page_count > MAX_PDF_PAGES:
            raise HTTPException(413, f"This PDF exceeds the {MAX_PDF_PAGES}-page limit.")
        outputs = []
        output_bytes = 0
        for index, page in enumerate(doc, start=1):
            pix = page.get_pixmap(dpi=dpi, alpha=False)
            dest = output_dir / f"page-{index:03d}.{ext}"
            if ext.lower() in {"jpg", "jpeg"}:
                pix.save(str(dest), jpg_quality=88)
            else:
                pix.save(str(dest))
            output_bytes += dest.stat().st_size
            if output_bytes > MAX_OUTPUT_BYTES:
                raise HTTPException(413, f"Rendered images exceed the {MAX_OUTPUT_BYTES // (1024 * 1024)} MB output limit. Reduce DPI or page count.")
            outputs.append(dest)
        doc.close()
        return outputs
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, "The PDF could not be rendered. It may be corrupted or use unsupported features.") from exc


def _images_to_pdf(images: list[Path], out: Path, page_size_name: str = "a4") -> Path:
    page_sizes = {"a4": (595.28, 841.89), "letter": (612.0, 792.0), "fit": None}
    page_size = page_sizes.get(page_size_name.lower(), page_sizes["a4"])
    try:
        dest = fitz.open()
        for image_path in images:
            with Image.open(image_path) as img:
                img = img.convert("RGB")
                width_px, height_px = img.size
            if page_size is None:
                width, height = width_px * 72 / 96, height_px * 72 / 96
            else:
                width, height = page_size
            page = dest.new_page(width=width, height=height)
            scale = min(width / width_px, height / height_px)
            iw, ih = width_px * scale, height_px * scale
            rect = fitz.Rect((width - iw) / 2, (height - ih) / 2, (width + iw) / 2, (height + ih) / 2)
            page.insert_image(rect, filename=str(image_path), keep_proportion=True)
        dest.save(str(out), garbage=4, deflate=True)
        dest.close()
        return out
    except Exception as exc:
        raise HTTPException(422, "One or more images could not be added to the PDF.") from exc


def _compare_pdfs(left: Path, right: Path, out: Path) -> Path:
    try:
        ldoc, rdoc = fitz.open(str(left)), fitz.open(str(right))
        page_count = max(ldoc.page_count, rdoc.page_count)
        page_size = landscape(A4)
        c = canvas.Canvas(str(out), pagesize=page_size)
        width, height = page_size
        c.setTitle("PDF comparison report")
        for i in range(page_count):
            c.setFillColor(colors.HexColor("#162033"))
            c.setFont("Helvetica-Bold", 16)
            c.drawString(28, height - 32, f"Page {i + 1} comparison")
            box_y, box_h = 45, height - 100
            box_w = (width - 84) / 2
            for j, doc in enumerate((ldoc, rdoc)):
                x = 28 + j * (box_w + 28)
                label = left.name if j == 0 else right.name
                c.setFont("Helvetica-Bold", 9)
                c.setFillColor(colors.HexColor("#46536A"))
                c.drawString(x, height - 52, label[:55])
                c.setStrokeColor(colors.HexColor("#D5DCE7"))
                c.roundRect(x, box_y, box_w, box_h, 6, stroke=1, fill=0)
                if i < doc.page_count:
                    pix = doc[i].get_pixmap(dpi=90, alpha=False)
                    img = ImageReader(io.BytesIO(pix.tobytes("png")))
                    iw, ih = img.getSize()
                    scale = min((box_w - 16) / iw, (box_h - 16) / ih)
                    dw, dh = iw * scale, ih * scale
                    c.drawImage(img, x + (box_w - dw) / 2, box_y + (box_h - dh) / 2, dw, dh, preserveAspectRatio=True, mask="auto")
                else:
                    c.setFont("Helvetica-Oblique", 12)
                    c.drawCentredString(x + box_w / 2, box_y + box_h / 2, "No corresponding page")
            c.setFont("Helvetica", 8)
            c.setFillColor(colors.HexColor("#657187"))
            c.drawString(28, 24, f"Source pages: {ldoc.page_count} | Comparison pages: {rdoc.page_count}. Raster previews are for visual review, not a text-level diff.")
            c.showPage()
        c.save()
        ldoc.close(); rdoc.close()
        return out
    except Exception as exc:
        raise HTTPException(422, "Could not compare these PDFs. Both files must be readable PDFs.") from exc


def _compare_pdfs_text(left: Path, right: Path, out: Path) -> Path:
    """Create a text-level, page-by-page comparison report; scanned pages are explicitly marked as unsearchable."""
    from xml.sax.saxutils import escape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.units import inch
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle

    try:
        ldoc, rdoc = fitz.open(str(left)), fitz.open(str(right))
        total = max(ldoc.page_count, rdoc.page_count)
        left_texts: list[str] = []
        right_texts: list[str] = []
        for idx in range(total):
            left_texts.append(ldoc[idx].get_text("text", sort=True) if idx < ldoc.page_count else "")
            right_texts.append(rdoc[idx].get_text("text", sort=True) if idx < rdoc.page_count else "")
        # Embed a Unicode-capable font when available so text comparisons can report Indic and accented text.
        regular_font, bold_font = "Helvetica", "Helvetica-Bold"
        regular_path = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
        bold_path = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")
        try:
            if regular_path.exists() and bold_path.exists():
                if "PaperNestDejaVu" not in pdfmetrics.getRegisteredFontNames():
                    pdfmetrics.registerFont(TTFont("PaperNestDejaVu", str(regular_path)))
                if "PaperNestDejaVu-Bold" not in pdfmetrics.getRegisteredFontNames():
                    pdfmetrics.registerFont(TTFont("PaperNestDejaVu-Bold", str(bold_path)))
                regular_font, bold_font = "PaperNestDejaVu", "PaperNestDejaVu-Bold"
        except Exception:
            regular_font, bold_font = "Helvetica", "Helvetica-Bold"
        styles = getSampleStyleSheet()
        styles.add(ParagraphStyle(name="DiffTitle", parent=styles["Title"], fontName=bold_font, fontSize=20, leading=25, textColor=colors.HexColor("#25334A"), alignment=TA_LEFT, spaceAfter=12))
        styles.add(ParagraphStyle(name="DiffHeading", parent=styles["Heading2"], fontName=bold_font, fontSize=13, leading=16, textColor=colors.HexColor("#2C3D58"), spaceBefore=9, spaceAfter=5))
        styles.add(ParagraphStyle(name="DiffBody", parent=styles["BodyText"], fontName=regular_font, fontSize=9, leading=12, spaceAfter=4, wordWrap="CJK"))
        styles.add(ParagraphStyle(name="DiffAdd", parent=styles["BodyText"], fontName=regular_font, fontSize=8, leading=10, textColor=colors.HexColor("#146C43"), backColor=colors.HexColor("#EAF7EF"), leftIndent=8, borderPadding=3, spaceAfter=3, wordWrap="CJK"))
        styles.add(ParagraphStyle(name="DiffRemove", parent=styles["BodyText"], fontName=regular_font, fontSize=8, leading=10, textColor=colors.HexColor("#A12E2E"), backColor=colors.HexColor("#FCECEC"), leftIndent=8, borderPadding=3, spaceAfter=3, wordWrap="CJK"))
        styles.add(ParagraphStyle(name="DiffMuted", parent=styles["BodyText"], fontName=regular_font, fontSize=8, leading=11, textColor=colors.HexColor("#667085"), spaceAfter=5))
        story: list[Any] = [Paragraph("PDF text comparison report", styles["DiffTitle"]), Paragraph(f"Left file: {escape(left.name)}<br/>Right file: {escape(right.name)}", styles["DiffBody"]), Spacer(1, 8)]
        changed_pages = 0
        only_left = max(0, ldoc.page_count - rdoc.page_count)
        only_right = max(0, rdoc.page_count - ldoc.page_count)
        total_diff_lines = 0
        entries_by_page = []
        for idx, (ltext, rtext) in enumerate(zip(left_texts, right_texts), start=1):
            lnorm = "\n".join(line.strip() for line in ltext.splitlines() if line.strip())
            rnorm = "\n".join(line.strip() for line in rtext.splitlines() if line.strip())
            if idx <= ldoc.page_count and idx <= rdoc.page_count and not lnorm and not rnorm:
                entries_by_page.append((idx, "empty", []))
                continue
            if lnorm == rnorm and idx <= ldoc.page_count and idx <= rdoc.page_count:
                entries_by_page.append((idx, "same", []))
                continue
            changed_pages += 1
            ll = lnorm.splitlines()[:1000]
            rr = rnorm.splitlines()[:1000]
            ops = difflib.SequenceMatcher(a=ll, b=rr, autojunk=True).get_opcodes()
            edits = []
            for tag, i1, i2, j1, j2 in ops:
                if tag == "equal":
                    continue
                for line in ll[i1:i2][:120]:
                    edits.append(("remove", line))
                for line in rr[j1:j2][:120]:
                    edits.append(("add", line))
                if len(edits) >= 240:
                    break
            total_diff_lines += len(edits)
            entries_by_page.append((idx, "diff", edits))
        summary_data = [
            ["Metric", "Result"],
            ["Left document pages", str(ldoc.page_count)],
            ["Right document pages", str(rdoc.page_count)],
            ["Pages with text differences or a missing counterpart", str(changed_pages)],
            ["Pages only in left document", str(only_left)],
            ["Pages only in right document", str(only_right)],
            ["Reported changed text lines", str(total_diff_lines)],
        ]
        table = Table(summary_data, colWidths=[3.5 * inch, 2.4 * inch], repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E9EEF5")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#25334A")),
            ("FONTNAME", (0, 0), (-1, 0), bold_font),
            ("FONTNAME", (0, 1), (-1, -1), regular_font),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), .45, colors.HexColor("#D4DCE7")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 7),
            ("RIGHTPADDING", (0, 0), (-1, -1), 7),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        story.extend([table, Spacer(1, 12), Paragraph("How to read this report", styles["DiffHeading"]), Paragraph("Green lines are additions in the right-hand PDF; red lines are lines found only in the left-hand PDF. Comparison is by extracted text lines at matching page numbers. Scanned pages without a text layer, complex reading order, and text encoded in unusual ways may not produce reliable text diffs. This report does not establish legal or factual equivalence.", styles["DiffMuted"]), PageBreak()])
        for idx, state, edits in entries_by_page:
            story.append(Paragraph(f"Page {idx}", styles["DiffHeading"]))
            ltext = left_texts[idx - 1]
            rtext = right_texts[idx - 1]
            if idx > ldoc.page_count:
                story.append(Paragraph("Page exists only in the right document.", styles["DiffMuted"]))
            elif idx > rdoc.page_count:
                story.append(Paragraph("Page exists only in the left document.", styles["DiffMuted"]))
            elif state == "same":
                story.append(Paragraph("No extracted text differences detected on this page.", styles["DiffMuted"]))
            elif state == "empty" or (not ltext.strip() and not rtext.strip()):
                story.append(Paragraph("No selectable text on either page. Text comparison is inconclusive; use visual comparison or OCR first.", styles["DiffMuted"]))
            else:
                if not edits:
                    story.append(Paragraph("Text differs but no line changes could be summarized (the diff may exceed the report cap).", styles["DiffMuted"]))
                for action, line in edits:
                    prefix = "+ " if action == "add" else "- "
                    safe_line = "".join(ch if ord(ch) >= 32 or ch in "\t" else " " for ch in line)
                    content = escape(prefix + (safe_line[:1200] + ("..." if len(safe_line) > 1200 else ""))) or "&#160;"
                    story.append(Paragraph(content, styles["DiffAdd"] if action == "add" else styles["DiffRemove"]))
                if len(ltext.splitlines()) > 1000 or len(rtext.splitlines()) > 1000 or len(edits) >= 240:
                    story.append(Paragraph("This page has more extracted text changes than the report cap; only an initial portion is shown.", styles["DiffMuted"]))
            story.append(Spacer(1, 7))
        SimpleDocTemplate(str(out), pagesize=A4, rightMargin=42, leftMargin=42, topMargin=42, bottomMargin=42, title="PDF text comparison report", author="PaperNest").build(story)
        ldoc.close(); rdoc.close()
        return out
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, "Could not create a text comparison report. Both PDFs must be readable and within processing limits.") from exc


def _form_field_details(path: Path) -> list[dict[str, Any]]:
    try:
        doc = fitz.open(str(path))
        if doc.needs_pass:
            doc.close()
            raise HTTPException(422, "This PDF form is password-protected. Unlock it with its known password before filling it.")
        if doc.page_count > MAX_PDF_PAGES:
            doc.close()
            raise HTTPException(413, f"Form inspection is limited to {MAX_PDF_PAGES} pages.")
        grouped: dict[str, list[tuple[Any, int]]] = {}
        supported = {"Text", "CheckBox", "RadioButton", "ComboBox", "ListBox"}
        for page_number, page in enumerate(doc, start=1):
            for widget in page.widgets() or []:
                name = str(widget.field_name or "").strip()
                kind = str(widget.field_type_string or "Unknown")
                if name and kind in supported:
                    grouped.setdefault(name, []).append((widget, page_number))
        if len(grouped) > 500:
            doc.close()
            raise HTTPException(413, "This form has more than 500 named fields. Simplify the form and retry.")
        details: list[dict[str, Any]] = []
        for name, entries in grouped.items():
            first_widget, page_number = entries[0]
            kind = str(first_widget.field_type_string or "Unknown")
            if kind == "RadioButton":
                choices = list(dict.fromkeys(str(widget.on_state()) for widget, _ in entries if widget.on_state() and str(widget.on_state()) != "Off"))
            else:
                choices = [str(value) for value in (first_widget.choice_values or [])]
            value: Any = first_widget.field_value
            if kind == "RadioButton":
                selected = [str(widget.on_state()) for widget, _ in entries if str(widget.field_value) == str(widget.on_state()) and str(widget.on_state()) != "Off"]
                value = selected[0] if selected else ""
            if not isinstance(value, (str, int, float, bool)):
                value = ""
            details.append({
                "name": name,
                "label": str(first_widget.field_label or name),
                "type": kind,
                "value": value,
                "options": choices,
                "page": page_number,
                "required": any(bool(int(widget.field_flags or 0) & 2) for widget, _ in entries),
            })
        doc.close()
        return details
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, "Could not inspect this PDF form. It may be damaged or use unsupported form features.") from exc


def _validate_public_http_url(url: str) -> str:
    """Reject local/private IPs and non-HTTP schemes before navigating or fetching resources."""
    try:
        parsed = urlsplit(url.strip())
        if parsed.scheme.lower() not in {"http", "https"}:
            raise HTTPException(400, "Enter a public HTTP or HTTPS URL.")
        if not parsed.hostname or parsed.username or parsed.password:
            raise HTTPException(400, "Enter a valid public URL without embedded credentials.")
        hostname = parsed.hostname.rstrip(".").lower()
        if hostname in {"localhost", "localhost.localdomain", "metadata.google.internal"} or hostname.endswith((".localhost", ".local", ".internal", ".lan")):
            raise HTTPException(400, "Private, local, and internal network URLs are not allowed.")
        try:
            ip = ipaddress.ip_address(hostname)
            addresses = [ip]
        except ValueError:
            try:
                infos = socket.getaddrinfo(hostname, parsed.port or (443 if parsed.scheme.lower() == "https" else 80), type=socket.SOCK_STREAM)
                addresses = list({ipaddress.ip_address(info[4][0].split("%", 1)[0]) for info in infos})
            except (OSError, ValueError) as exc:
                raise HTTPException(400, "The URL hostname could not be resolved. Check the URL and try again.") from exc
        if not addresses or any(not address.is_global for address in addresses):
            raise HTTPException(400, "The URL must resolve only to publicly routable IP addresses. Private and local network targets are blocked.")
        return url.strip()
    except HTTPException:
        raise
    except ValueError as exc:
        raise HTTPException(400, "The URL or port is invalid.") from exc


def _html_to_pdf(paths: list[Path], options: dict[str, Any], outdir: Path, base: str) -> tuple[Path, str]:
    """Render a public HTTP(S) page or a user-uploaded HTML file in a temporary browser."""
    url = str(options.get("url", "")).strip()
    source_html = ""
    if paths and url:
        raise HTTPException(400, "Choose either a webpage URL or an uploaded HTML file, not both.")
    if paths:
        if len(paths) != 1 or paths[0].suffix.lower() not in HTML_EXTS:
            raise HTTPException(400, "Upload one HTML file, or leave files empty and enter a public URL.")
        if paths[0].stat().st_size > 2 * 1024 * 1024:
            raise HTTPException(413, "Uploaded HTML is limited to 2 MB.")
        source_html = paths[0].read_text(encoding="utf-8", errors="replace")
        if not source_html.strip():
            raise HTTPException(400, "The uploaded HTML file is empty.")
    elif not url:
        raise HTTPException(400, "Enter a public URL or upload an HTML file.")
    if url:
        url = _validate_public_http_url(url)

    try:
        from playwright.sync_api import TimeoutError as PlaywrightTimeoutError
        from playwright.sync_api import sync_playwright
    except ImportError as exc:
        raise HTTPException(503, "HTML-to-PDF requires Playwright and Chromium. Install the documented browser dependencies and restart the API.") from exc

    outdir.mkdir(parents=True, exist_ok=True)
    target = outdir / f"{safe_stem(base)}.pdf"
    blocked_requests: list[str] = []
    request_count = 0
    max_requests = 200

    def browser_request_guard(route: Any) -> None:
        nonlocal request_count
        request_count += 1
        if request_count > max_requests:
            blocked_requests.append("request limit")
            route.abort()
            return
        request_url = route.request.url
        if urlsplit(request_url).scheme.lower() == "data":
            route.continue_()
            return
        try:
            _validate_public_http_url(request_url)
            route.continue_()
        except Exception:
            blocked_requests.append(request_url[:180])
            route.abort()

    try:
        with sync_playwright() as playwright:
            launch_options: dict[str, Any] = {
                "headless": True,
                "args": ["--disable-dev-shm-usage", "--disable-background-networking", "--no-first-run", "--disable-extensions"],
            }
            chromium_path = shutil.which("chromium") or shutil.which("chromium-browser")
            if chromium_path:
                launch_options["executable_path"] = chromium_path
            browser = playwright.chromium.launch(**launch_options)
            context = browser.new_context(
                java_script_enabled=bool(url),
                accept_downloads=False,
                service_workers="block",
                viewport={"width": 1280, "height": 900},
                device_scale_factor=1,
            )
            page = context.new_page()
            page.route("**/*", browser_request_guard)
            if url:
                try:
                    page.goto(url, wait_until="domcontentloaded", timeout=30000)
                    try:
                        page.wait_for_load_state("networkidle", timeout=7000)
                    except PlaywrightTimeoutError:
                        pass
                    page.wait_for_timeout(400)
                except PlaywrightTimeoutError as exc:
                    raise HTTPException(504, "The webpage took too long to load. Try a simpler page or upload its HTML directly.") from exc
            else:
                # JavaScript stays disabled for uploaded HTML to reduce script-based abuse.
                page.set_content(source_html, wait_until="domcontentloaded", timeout=10000)
                try:
                    page.wait_for_load_state("networkidle", timeout=3000)
                except PlaywrightTimeoutError:
                    pass

            page_format_key = str(options.get("page_format", "A4")).upper()
            page_format = {"A4": "A4", "LETTER": "Letter", "LEGAL": "Legal", "A3": "A3", "A5": "A5"}.get(page_format_key, "A4")
            margin_mm = max(0, min(30, float(options.get("margin_mm", 10))))
            print_background = str(options.get("print_background", "yes")).lower() in {"yes", "true", "1"}
            page.pdf(
                path=str(target),
                format=page_format,
                print_background=print_background,
                prefer_css_page_size=True,
                margin={"top": f"{margin_mm}mm", "right": f"{margin_mm}mm", "bottom": f"{margin_mm}mm", "left": f"{margin_mm}mm"},
                display_header_footer=False,
            )
            context.close()
            browser.close()
    except HTTPException:
        target.unlink(missing_ok=True)
        raise
    except Exception as exc:
        target.unlink(missing_ok=True)
        message = "Could not render this webpage to PDF. Confirm Chromium is installed and try a simpler page."
        if "Executable doesn't exist" in str(exc) or "browserType.launch" in str(exc):
            message = "Chromium is not installed for the HTML renderer. Install Chromium using the deployment instructions and restart the API."
        raise HTTPException(422, message) from exc

    if not target.exists() or target.stat().st_size == 0:
        raise HTTPException(422, "The webpage renderer did not produce a PDF.")
    if target.stat().st_size > MAX_OUTPUT_BYTES:
        target.unlink(missing_ok=True)
        raise HTTPException(413, f"The generated PDF exceeds the {MAX_OUTPUT_BYTES // (1024 * 1024)} MB output limit.")
    try:
        rendered = fitz.open(str(target))
        page_count = rendered.page_count
        rendered.close()
    except Exception as exc:
        target.unlink(missing_ok=True)
        raise HTTPException(422, "The HTML renderer generated an invalid PDF. Try a different page.") from exc
    max_html_pages = min(200, MAX_PDF_PAGES)
    if page_count < 1 or page_count > max_html_pages:
        target.unlink(missing_ok=True)
        raise HTTPException(413, f"HTML-to-PDF output is limited to {max_html_pages} pages per job; this page rendered to {page_count} pages.")
    source_label = urlsplit(url).hostname if url else f"{base}.html"
    warning = " External resources pointing to private/local network targets are blocked." if blocked_requests else ""
    js_note = " JavaScript is enabled for public URLs and disabled for uploaded HTML files."
    return target, f"Converted {source_label} to a {page_format} PDF with {page_count} page(s). Public URL rendering is network-filtered.{warning}{js_note}"


def _process_single(tool: str, paths: list[Path], options: dict[str, Any], outdir: Path, base: str) -> tuple[Path, str]:
    outdir.mkdir(parents=True, exist_ok=True)
    if tool == "html-to-pdf":
        return _html_to_pdf(paths, options, outdir, base)
    if tool == "merge-pdf":
        if len(paths) < 2:
            raise HTTPException(400, "Select at least two PDFs to merge.")
        writer = PdfWriter()
        for path in paths:
            reader = _read_pdf(path)
            for page in reader.pages:
                writer.add_page(page)
        return _write_pdf(writer, outdir / f"{base}-merged.pdf"), f"Merged {len(paths)} PDF files."

    if tool in {"jpg-to-pdf", "png-to-pdf"}:
        expected = {"jpg-to-pdf": {".jpg", ".jpeg"}, "png-to-pdf": {".png"}}[tool]
        if any(p.suffix.lower() not in expected for p in paths):
            ext_text = "JPG/JPEG" if tool == "jpg-to-pdf" else "PNG"
            raise HTTPException(400, f"{tool} accepts {ext_text} images only.")
        return _images_to_pdf(paths, outdir / f"{base}.pdf", str(options.get("page_size", "a4"))), f"Created a PDF with {len(paths)} image page(s)."

    if tool == "compare-pdf":
        if len(paths) != 2:
            raise HTTPException(400, "Upload exactly two PDFs to compare them.")
        for p in paths:
            _read_pdf(p)
        mode = str(options.get("mode", "both"))
        if mode not in {"visual", "text", "both"}:
            raise HTTPException(400, "Comparison mode must be visual, text, or both.")
        target = outdir / f"{base}-comparison.pdf"
        if mode == "visual":
            return _compare_pdfs(paths[0], paths[1], target), "Created a side-by-side visual comparison report. The report compares page appearance, not extracted text."
        if mode == "text":
            return _compare_pdfs_text(paths[0], paths[1], target), "Created a text-level comparison report with changed pages and line-level additions/removals. Scanned pages without a text layer cannot be compared semantically."
        visual_path = _compare_pdfs(paths[0], paths[1], outdir / f"{base}-visual.pdf")
        text_path = _compare_pdfs_text(paths[0], paths[1], outdir / f"{base}-text-diff.pdf")
        writer = PdfWriter()
        report_readers = [PdfReader(str(report_path), strict=False) for report_path in (visual_path, text_path)]
        for report_reader in report_readers:
            for page in report_reader.pages:
                writer.add_page(page)
        _write_pdf(writer, target)
        return target, "Created a combined report: side-by-side page previews followed by a text-level comparison. Text comparison is limited for scanned pages, complex reading order, and visually identical text represented differently."

    if tool == "redact-pdf":
        if len(paths) != 1 or paths[0].suffix.lower() != ".pdf":
            raise HTTPException(400, "Upload exactly one PDF to redact.")
        if str(options.get("confirmed", "no")).lower() not in {"yes", "true", "1"}:
            raise HTTPException(400, "Confirm that you understand redaction permanently removes content before continuing.")
        terms_raw = str(options.get("terms", ""))[:10000]
        terms = list(dict.fromkeys(line.strip() for line in terms_raw.splitlines() if line.strip()))
        boxes_raw = str(options.get("boxes_json", "")).strip()
        try:
            boxes = json.loads(boxes_raw) if boxes_raw else []
        except json.JSONDecodeError as exc:
            raise HTTPException(400, "Page regions must be valid JSON: an array of page/x0/y0/x1/y1 objects.") from exc
        if not isinstance(boxes, list) or len(boxes) > 200:
            raise HTTPException(400, "Page regions must be a JSON array with no more than 200 rectangles.")
        if not terms and not boxes:
            raise HTTPException(400, "Enter at least one text phrase or page region to redact.")
        if len(terms) > 100:
            raise HTTPException(400, "Use no more than 100 text phrases per redaction job.")
        if sum(len(term) for term in terms) > 5000:
            raise HTTPException(400, "Redaction phrases are too long. Use shorter phrases.")
        pages_spec = str(options.get("pages", "")).strip()
        fill_mode = str(options.get("fill", "black")).lower()
        if fill_mode not in {"black", "white", "blur"}:
            raise HTTPException(400, "Redaction style must be black, white, or blur.")
        # Blur mode first permanently removes the original content, then overlays a blurred
        # raster copy of the selected region. This avoids leaving searchable source text underneath.
        fill = (1, 1, 1) if fill_mode in {"white", "blur"} else (0, 0, 0)
        try:
            doc = fitz.open(str(paths[0]))
            page_count = doc.page_count
            if page_count > MAX_PDF_PAGES:
                doc.close()
                raise HTTPException(413, f"Redaction is limited to {MAX_PDF_PAGES} pages per file.")
            selected = _parse_pages(pages_spec, page_count, default_all=True) if pages_spec else list(range(page_count))
            selected_set = set(selected)
            planned = 0
            pages_with_redactions: set[int] = set()
            blur_regions: list[tuple[int, fitz.Rect, bytes]] = []

            def capture_blur_patch(page: fitz.Page, page_index: int, rect: fitz.Rect) -> None:
                # Render only the marked region before redaction, blur it, and store only the
                # blurred pixels for reinsertion after the underlying PDF content is removed.
                clip = rect & page.rect
                if clip.is_empty or clip.is_infinite:
                    return
                pix = page.get_pixmap(matrix=fitz.Matrix(3, 3), clip=clip, alpha=False, annots=False)
                with Image.open(io.BytesIO(pix.tobytes("png"))) as source_image:
                    blurred = source_image.convert("RGB").filter(ImageFilter.GaussianBlur(radius=10))
                    encoded = io.BytesIO()
                    blurred.save(encoded, format="PNG", optimize=True)
                    blur_regions.append((page_index, fitz.Rect(clip), encoded.getvalue()))

            for page_index in selected:
                page = doc[page_index]
                for term in terms:
                    for quad in page.search_for(term, quads=True):
                        quad_rect = quad.rect
                        if fill_mode == "blur":
                            capture_blur_patch(page, page_index, quad_rect)
                        page.add_redact_annot(quad, fill=fill, cross_out=False)
                        pages_with_redactions.add(page_index)
                        planned += 1
            for item in boxes:
                if not isinstance(item, dict):
                    raise HTTPException(400, "Each page region must be an object with page, x0, y0, x1, and y1 values.")
                try:
                    page_number = int(item.get("page", 0))
                    x0, y0, x1, y1 = (float(item[key]) for key in ("x0", "y0", "x1", "y1"))
                except (TypeError, ValueError, KeyError) as exc:
                    raise HTTPException(400, "Each page region must include numeric page, x0, y0, x1, and y1 values.") from exc
                if page_number < 1 or page_number > page_count:
                    raise HTTPException(400, f"Redaction region page {page_number} is outside this PDF's {page_count} pages.")
                page_index = page_number - 1
                if page_index not in selected_set:
                    raise HTTPException(400, f"Page region {page_number} is outside the selected page range.")
                page = doc[page_index]
                rect = fitz.Rect(x0, y0, x1, y1)
                if not all(map(lambda v: abs(v) < 1e7, (x0, y0, x1, y1))) or rect.is_empty or rect.is_infinite or not page.rect.contains(rect):
                    raise HTTPException(400, f"Redaction rectangle on page {page_number} must be non-empty and within the visible page in PDF points.")
                # The visual selector maps coordinates from the page as rendered in the viewer.
                # For pages with /Rotate set, PyMuPDF drawing/edit coordinates are in unrotated space.
                if bool(item.get("visual", False)):
                    rect = rect * page.derotation_matrix
                    if rect.is_empty or rect.is_infinite:
                        raise HTTPException(400, f"Redaction region on page {page_number} could not be mapped to page coordinates.")
                if fill_mode == "blur":
                    capture_blur_patch(page, page_index, rect)
                page.add_redact_annot(rect, fill=fill, cross_out=False)
                pages_with_redactions.add(page_index)
                planned += 1
            if planned == 0:
                doc.close()
                raise HTTPException(422, "No text matches or redaction regions were found. No output was created.")
            for page_index in sorted(pages_with_redactions):
                page = doc[page_index]
                page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_PIXELS, graphics=fitz.PDF_REDACT_LINE_ART_REMOVE_IF_COVERED, text=fitz.PDF_REDACT_TEXT_REMOVE)
            if fill_mode == "blur":
                for page_index, rect, patch in blur_regions:
                    # The blurred image is the only remaining representation of the marked content.
                    doc[page_index].insert_image(rect, stream=patch, overlay=True, keep_proportion=False)
            # Redaction is a sanitization operation: also remove potential copies of sensitive content in metadata,
            # document-level metadata, attachments, comments, links, and form widgets.
            try:
                doc.set_metadata({})
            except Exception:
                pass
            try:
                doc.del_xml_metadata()
            except Exception:
                pass
            try:
                doc.set_toc([])
            except Exception:
                pass
            for page in doc:
                for widget in list(page.widgets() or []):
                    try:
                        page.delete_widget(widget)
                    except Exception:
                        pass
                for annot in list(page.annots() or []):
                    try:
                        page.delete_annot(annot)
                    except Exception:
                        pass
                for link in list(page.get_links() or []):
                    try:
                        page.delete_link(link)
                    except Exception:
                        pass
            for name in list(doc.embfile_names() or []):
                try:
                    doc.embfile_del(name)
                except Exception:
                    pass
            target = outdir / f"{base}-redacted.pdf"
            doc.save(str(target), garbage=4, deflate=True, clean=True)
            doc.close()
            # Verify that searched phrases are no longer extractable on the selected pages after saving.
            with fitz.open(str(target)) as check_doc:
                residual = []
                for page_index in selected:
                    text = check_doc[page_index].get_text("text")
                    folded = text.casefold()
                    for term in terms:
                        if term.casefold() in folded:
                            residual.append(term)
                if residual:
                    target.unlink(missing_ok=True)
                    raise HTTPException(422, "Output verification found a redaction phrase still present in page text. No output was retained.")
                if check_doc.page_count != page_count:
                    target.unlink(missing_ok=True)
                    raise HTTPException(422, "Output verification failed because the page count changed unexpectedly.")
            style_label = {"black": "black-filled", "white": "white-filled", "blur": "blurred"}[fill_mode]
            return target, f"Applied and verified {planned} {style_label} redaction region(s) across {len(selected)} selected page(s). Underlying text and objects were removed before the selected appearance was applied; metadata, attachments, comments, links, and form fields were also removed for sanitization. Redaction is irreversible. Review the output before sharing it."
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(422, "Could not redact this PDF. It may be encrypted, damaged, or contain unsupported content.") from exc

    if tool == "pdf-reader":
        if len(paths) != 1:
            raise HTTPException(400, "Open one PDF at a time in the reader.")
        reader = _read_pdf(paths[0])
        out = outdir / f"{base}-reader-copy.pdf"
        shutil.copyfile(paths[0], out)
        return out, f"PDF opened successfully ({len(reader.pages)} pages)."

    # All remaining operations start from a single source document. A second image is allowed for image watermarking / editing.
    if not paths:
        raise HTTPException(400, "Upload a file to continue.")
    source = paths[0]
    ext = source.suffix.lower()

    if tool == "sanitize-pdf-metadata":
        if ext != ".pdf":
            raise HTTPException(400, "Metadata sanitization requires a PDF file.")
        scope = str(options.get("scope", "metadata")).lower()
        if scope not in {"metadata", "full"}:
            raise HTTPException(400, "Sanitization scope must be metadata or full.")
        target = outdir / f"{base}-sanitized.pdf"
        try:
            doc = fitz.open(str(source))
            if doc.is_encrypted:
                doc.close()
                raise HTTPException(422, "This PDF is password-protected. Unlock it before sanitizing metadata.")
            if scope == "metadata":
                doc.scrub(attached_files=False, clean_pages=False, embedded_files=False, hidden_text=False, javascript=False, metadata=True, redactions=False, redact_images=0, remove_links=False, reset_fields=False, reset_responses=False, thumbnails=True, xml_metadata=True)
            else:
                doc.scrub(attached_files=True, clean_pages=True, embedded_files=True, hidden_text=True, javascript=True, metadata=True, redactions=False, redact_images=0, remove_links=True, reset_fields=True, reset_responses=True, thumbnails=True, xml_metadata=True)
                for page in doc:
                    for annot in list(page.annots() or []):
                        page.delete_annot(annot)
                    for widget in list(page.widgets() or []):
                        page.delete_widget(widget)
                doc.set_toc([])
            doc.set_metadata({})
            try:
                doc.del_xml_metadata()
            except Exception:
                pass
            doc.save(str(target), garbage=4, deflate=True, clean=True)
            doc.close()
            verify = fitz.open(str(target))
            metadata = verify.metadata or {}
            residual = [value for key, value in metadata.items() if key not in {"format", "encryption"} and value not in (None, "")]
            try:
                xmp = verify.get_xml_metadata() or ""
            except Exception:
                xmp = ""
            if residual or xmp.strip():
                verify.close()
                raise HTTPException(422, "Metadata verification found residual document metadata. No output was kept.")
            page_count = verify.page_count
            verify.close()
            if scope == "metadata":
                return target, f"Removed and verified standard and XMP metadata from the {page_count}-page PDF. Links, annotations, form fields, and attachments were left intact."
            return target, f"Full sanitization removed standard/XMP metadata and attempted to remove attachments, JavaScript, links, comments/annotations, form widgets, thumbnails, hidden text, and document outlines across {page_count} page(s). Review the document because some interactive features may be removed."
        except HTTPException:
            target.unlink(missing_ok=True)
            raise
        except Exception as exc:
            target.unlink(missing_ok=True)
            raise HTTPException(422, "Could not sanitize this PDF. It may be encrypted, malformed, or use unsupported features.") from exc

    if tool == "word-to-pdf":
        if ext not in WORD_EXTS:
            raise HTTPException(400, "Upload a DOC or DOCX file.")
        result = _run_libreoffice(source, outdir)
        target = outdir / f"{base}.pdf"
        if result != target:
            shutil.copyfile(result, target)
        return target, "Converted the Word document to PDF with LibreOffice. Layout may differ from Word depending on available fonts and settings."

    if tool == "excel-to-pdf":
        if ext not in EXCEL_EXTS:
            raise HTTPException(400, "Upload an XLS or XLSX file.")
        result = _run_libreoffice(source, outdir)
        target = outdir / f"{base}.pdf"
        if result != target:
            shutil.copyfile(result, target)
        return target, "Converted the spreadsheet to PDF with LibreOffice. Verify print areas, page scaling, and multi-sheet output."

    if tool == "powerpoint-to-pdf":
        if ext not in PPT_EXTS:
            raise HTTPException(400, "Upload a PPT or PPTX file.")
        result = _run_libreoffice(source, outdir)
        target = outdir / f"{base}.pdf"
        if result != target:
            shutil.copyfile(result, target)
        return target, "Converted the presentation to PDF with LibreOffice. Verify fonts and slide layout."

    if tool == "protect-pdf":
        if len(paths) != 1 or paths[0].suffix.lower() != ".pdf":
            raise HTTPException(400, "Upload one unprotected PDF to protect it.")
        password = str(options.get("password", ""))
        confirmation = str(options.get("password_confirm", ""))
        if len(password) < 8 or not password.strip():
            raise HTTPException(400, "Choose a non-empty password with at least 8 characters.")
        if len(password) > 128:
            raise HTTPException(400, "Password must be 128 characters or fewer.")
        if password != confirmation:
            raise HTTPException(400, "The password and confirmation do not match.")
        try:
            reader = PdfReader(str(paths[0]), strict=False)
            if reader.is_encrypted:
                raise HTTPException(422, "This PDF is already password-protected. Unlock it first if you know the password.")
            writer = PdfWriter()
            for page in reader.pages:
                writer.add_page(page)
            if reader.metadata:
                metadata = {str(k): str(v) for k, v in reader.metadata.items() if k and v is not None}
                if metadata:
                    writer.add_metadata(metadata)
            writer.encrypt(password, algorithm="AES-256")
            target = outdir / f"{base}-protected.pdf"
            _write_pdf(writer, target)
            check = PdfReader(str(target), strict=False)
            if not check.is_encrypted:
                raise HTTPException(500, "The PDF engine did not apply encryption to the output.")
            return target, "Password protection applied using AES-256 encryption. Keep your password safe; it cannot be recovered by PaperNest."
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(422, "Could not protect this PDF. Verify that it is a readable, unencrypted document.") from exc

    if tool == "unlock-pdf":
        if len(paths) != 1 or paths[0].suffix.lower() != ".pdf":
            raise HTTPException(400, "Upload one password-protected PDF to unlock it.")
        password = str(options.get("password", ""))
        if not password:
            raise HTTPException(400, "Enter the current PDF password.")
        try:
            reader = PdfReader(str(paths[0]), strict=False)
            if reader.is_encrypted:
                try:
                    decrypted = reader.decrypt(password)
                except Exception:
                    decrypted = 0
                if not decrypted:
                    raise HTTPException(400, "That password did not unlock the PDF. Check it and try again.")
            writer = PdfWriter()
            for page in reader.pages:
                writer.add_page(page)
            if reader.metadata:
                metadata = {str(k): str(v) for k, v in reader.metadata.items() if k and v is not None}
                if metadata:
                    writer.add_metadata(metadata)
            target = outdir / f"{base}-unlocked.pdf"
            _write_pdf(writer, target)
            check = PdfReader(str(target), strict=False)
            if check.is_encrypted:
                raise HTTPException(500, "The generated PDF is still encrypted.")
            return target, "Created an unencrypted copy using the password you provided. This does not bypass unknown passwords."
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(422, "Could not unlock this PDF. It may use an unsupported encryption mode or be damaged.") from exc

    if tool == "fill-pdf-form":
        if len(paths) != 1 or paths[0].suffix.lower() != ".pdf":
            raise HTTPException(400, "Upload exactly one PDF form to fill.")
        values = options.get("form_values", {})
        if not isinstance(values, dict):
            raise HTTPException(400, "Form values must be a JSON object keyed by field name.")
        values = {str(key): value for key, value in values.items() if value not in (None, "")}
        if not values:
            raise HTTPException(400, "Enter at least one form field value before processing.")
        try:
            doc = fitz.open(str(paths[0]))
            if doc.page_count > MAX_PDF_PAGES:
                doc.close()
                raise HTTPException(413, f"Form filling is limited to {MAX_PDF_PAGES} pages.")
            field_map = {}
            for page_number, page in enumerate(doc, start=1):
                for widget in list(page.widgets() or []):
                    name = str(widget.field_name or "")
                    if name:
                        field_map.setdefault(name, []).append((page, widget, page_number))
            if not field_map:
                doc.close()
                raise HTTPException(422, "No fillable AcroForm fields were found. Scanned and flat PDFs need a separate text-stamping workflow.")
            unknown = sorted(set(values) - set(field_map))
            if unknown:
                doc.close()
                raise HTTPException(400, "Unknown form field name(s): " + ", ".join(unknown[:10]))
            changed = 0
            for name, raw_value in values.items():
                entries = field_map[name]
                kind = str(entries[0][1].field_type_string or "")
                if kind in {"Button", "Signature"}:
                    raise HTTPException(400, f"Field '{name}' is a button or signature field and cannot be filled by this tool.")
                if kind == "RadioButton":
                    requested = str(raw_value)
                    choices = list(dict.fromkeys(str(entry[1].on_state()) for entry in entries if entry[1].on_state() and str(entry[1].on_state()) != "Off"))
                    if choices and requested not in choices:
                        raise HTTPException(400, f"Invalid option for radio field '{name}'. Choose one of: {', '.join(choices)}.")
                    for _, radio_widget, _ in entries:
                        radio_widget.field_value = requested if str(radio_widget.on_state()) == requested else "Off"
                        try:
                            radio_widget.update()
                        except Exception as exc:
                            raise HTTPException(422, f"Could not update radio field '{name}'.") from exc
                        changed += 1
                    continue
                for page, widget, page_number in entries:
                    if kind == "CheckBox":
                        if isinstance(raw_value, bool):
                            widget.field_value = (widget.on_state() or "Yes") if raw_value else "Off"
                        else:
                            truthy = str(raw_value).strip().lower() in {"1", "true", "yes", "on", str(widget.on_state() or "yes").lower()}
                            widget.field_value = (widget.on_state() or "Yes") if truthy else "Off"
                    elif kind in {"ComboBox", "ListBox"}:
                        choice_values = [str(v) for v in (widget.choice_values or [])]
                        requested_values = raw_value if isinstance(raw_value, list) else [raw_value]
                        requested = [str(v) for v in requested_values]
                        if choice_values and any(v not in choice_values for v in requested):
                            raise HTTPException(400, f"Invalid option for field '{name}'. Choose from the options displayed by the PDF form.")
                        widget.field_value = requested if kind == "ListBox" and isinstance(raw_value, list) else (requested[0] if requested else "")
                    else:
                        widget.field_value = str(raw_value)[:5000]
                    try:
                        widget.update()
                    except Exception as exc:
                        raise HTTPException(422, f"Could not update form field '{name}'. The field may use an unsupported font or field type.") from exc
                    changed += 1
            flatten = str(options.get("flatten", "no")).lower() in {"yes", "true", "1"}
            if flatten:
                doc.bake(annots=False, widgets=True)
            target = outdir / f"{base}-filled-form.pdf"
            doc.save(str(target), garbage=4, deflate=True, clean=True)
            doc.close()
            with fitz.open(str(target)) as verification:
                remaining = sum(1 for page in verification for _ in (page.widgets() or []))
                if flatten and remaining:
                    target.unlink(missing_ok=True)
                    raise HTTPException(422, "Output verification failed: some form fields remain interactive after flattening.")
                if not flatten and remaining < 1:
                    target.unlink(missing_ok=True)
                    raise HTTPException(422, "Output verification failed: expected editable form fields were not preserved.")
            return target, f"Filled {len(values)} form field name(s) across {changed} widget(s). " + ("Form fields were flattened into page content." if flatten else "Form fields remain editable.")
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(422, "Could not fill this form. The PDF may be damaged or use unsupported form features.") from exc

    if ext != ".pdf":
        raise HTTPException(400, "This tool requires a PDF file as its first upload.")

    reader = _read_pdf(source)
    total = len(reader.pages)

    if tool == "ocr-pdf":
        if source.suffix.lower() != ".pdf":
            raise HTTPException(400, "Upload a PDF file.")
        executable = shutil.which("tesseract")
        if not executable:
            raise HTTPException(503, "OCR requires Tesseract OCR. Install Tesseract and its language data on the API host.")
        language = str(options.get("language", "eng"))
        if not re.fullmatch(r"[A-Za-z0-9_+]+", language):
            raise HTTPException(400, "Choose a valid OCR language code.")
        try:
            available = subprocess.run([executable, "--list-langs"], capture_output=True, text=True, timeout=10, check=False)
            installed = {line.strip() for line in available.stdout.splitlines()[1:] if line.strip()}
            requested = language.split("+")
            if not all(code in installed for code in requested):
                missing = ", ".join(code for code in requested if code not in installed)
                raise HTTPException(400, f"OCR language data is not installed: {missing}. Install the matching Tesseract language package.")
            doc = fitz.open(str(source))
            if doc.page_count > MAX_OCR_PAGES:
                doc.close()
                raise HTTPException(413, f"OCR is limited to {MAX_OCR_PAGES} pages per file. Split the PDF and retry.")
            writer = PdfWriter()
            rendered_total = 0
            for index, page in enumerate(doc, start=1):
                page_dir = outdir / f"ocr-page-{index:03d}"
                page_dir.mkdir(parents=True, exist_ok=True)
                pix = page.get_pixmap(dpi=200, alpha=False)
                image_path = page_dir / "page.png"
                pix.save(str(image_path))
                out_base = page_dir / "recognized"
                try:
                    run = subprocess.run([executable, str(image_path), str(out_base), "-l", language, "pdf"], capture_output=True, timeout=90, check=False)
                except subprocess.TimeoutExpired as exc:
                    doc.close()
                    raise HTTPException(504, f"OCR timed out while processing page {index}. Try fewer pages or a smaller PDF.") from exc
                page_pdf = page_dir / "recognized.pdf"
                if run.returncode != 0 or not page_pdf.exists() or page_pdf.stat().st_size == 0:
                    doc.close()
                    raise HTTPException(422, f"OCR failed on page {index}. Check the selected language and PDF page quality.")
                rendered_total += page_pdf.stat().st_size
                if rendered_total > MAX_OUTPUT_BYTES:
                    doc.close()
                    raise HTTPException(413, f"OCR output exceeds the {MAX_OUTPUT_BYTES // (1024 * 1024)} MB output limit.")
                recognized_reader = PdfReader(str(page_pdf), strict=False)
                for recognized_page in recognized_reader.pages:
                    writer.add_page(recognized_page)
            doc.close()
            target = outdir / f"{base}-ocr.pdf"
            _write_pdf(writer, target)
            return target, f"Created a searchable OCR PDF for {writer.get_num_pages()} page(s) using language '{language}'. OCR quality depends on scan resolution and text quality; review the result before relying on it."
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(422, "Could not OCR this PDF. It may be damaged or contain unsupported page content.") from exc

    if tool == "pdf-to-text":
        if source.suffix.lower() != ".pdf":
            raise HTTPException(400, "Upload a PDF file.")
        try:
            doc = fitz.open(str(source))
            parts = []
            has_selectable_text = False
            for index, page in enumerate(doc, start=1):
                page_text = page.get_text("text", sort=True).strip()
                if page_text:
                    has_selectable_text = True
                parts.append(f"===== Page {index} =====\n{page_text}")
            doc.close()
            extracted = "\n\n".join(parts).strip()
            if not has_selectable_text:
                raise HTTPException(422, "No selectable text was found. Scanned PDFs need OCR before text can be extracted.")
            target = outdir / f"{base}.txt"
            target.write_text(extracted + "\n", encoding="utf-8")
            return target, "Extracted selectable PDF text into UTF-8 TXT, separated by page. Scanned pages are not OCR-processed."
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(422, "Could not extract text from this PDF.") from exc

    if tool == "pdf-metadata":
        try:
            doc = fitz.open(str(source))
            current = doc.metadata or {}
            fields = ("title", "author", "subject", "keywords", "creator")
            updated = dict(current)
            changed = []
            for field in fields:
                value = str(options.get(field, ""))[:1000].strip()
                if value:
                    updated[field] = value
                    changed.append(field)
            if not changed:
                doc.close()
                raise HTTPException(400, "Enter at least one metadata field to update.")
            doc.set_metadata(updated)
            target = outdir / f"{base}-metadata.pdf"
            doc.save(str(target), garbage=4, deflate=True)
            doc.close()
            return target, "Updated PDF document metadata: " + ", ".join(changed) + ". Empty fields were left unchanged."
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(422, "Could not update the PDF metadata.") from exc

    if tool == "pdf-to-png":
        dpi = min(240, max(72, int(options.get("dpi", 144))))
        renders = _render_pdf_pages(source, outdir / "png-pages", dpi=dpi, ext="png")
        return _make_zip([(p, p.name) for p in renders], outdir / f"{base}-png-pages.zip"), f"Rendered {len(renders)} page(s) as PNG images at {dpi} DPI."

    if tool == "sign-pdf":
        if len(paths) != 2 or paths[0].suffix.lower() != ".pdf" or paths[1].suffix.lower() not in IMAGE_EXTS:
            raise HTTPException(400, "Upload one PDF first, then one signature image (PNG with transparency recommended, or JPG).")
        try:
            doc = fitz.open(str(source))
            page_count = doc.page_count
            chosen = _parse_pages(str(options.get("pages", "")), page_count, default_all=False) if str(options.get("pages", "")).strip() else [page_count - 1]
            x_pct = max(0, min(100, float(options.get("x", 65))))
            y_pct = max(0, min(100, float(options.get("y", 82))))
            width_pct = max(2, min(70, float(options.get("width", 25))))
            height_pct = max(2, min(40, float(options.get("height", 10))))
            for index in chosen:
                page = doc[index]
                left = min(page.rect.width - 2, page.rect.width * x_pct / 100)
                top = min(page.rect.height - 2, page.rect.height * y_pct / 100)
                width = page.rect.width * width_pct / 100
                height = page.rect.height * height_pct / 100
                rect = fitz.Rect(left, top, min(page.rect.width - 1, left + width), min(page.rect.height - 1, top + height))
                page.insert_image(rect, filename=str(paths[1]), keep_proportion=True, overlay=True)
            target = outdir / f"{base}-signed-copy.pdf"
            doc.save(str(target), garbage=4, deflate=True)
            doc.close()
            return target, f"Placed the signature image on {len(chosen)} page(s). This is a visual signature overlay, not a cryptographic digital signature."
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(422, "Could not place the signature image on this PDF.") from exc

    if tool == "split-pdf":
        spec = str(options.get("pages", "")).strip()
        mode = str(options.get("mode", "ranges"))
        selected = _parse_pages(spec, total) if spec else list(range(total))
        if not spec or mode == "each":
            outputs = []
            for page_no in selected:
                writer = PdfWriter(); writer.add_page(reader.pages[page_no])
                part = outdir / f"{base}-page-{page_no + 1:03d}.pdf"
                _write_pdf(writer, part); outputs.append((part, part.name))
            return _make_zip(outputs, outdir / f"{base}-split.zip"), f"Split the selection into {len(outputs)} one-page PDF file(s)."
        if mode == "selected":
            writer = PdfWriter()
            for page_no in selected:
                writer.add_page(reader.pages[page_no])
            return _write_pdf(writer, outdir / f"{base}-selected-pages.pdf"), f"Combined {len(selected)} selected page(s) into one PDF."
        outputs = []
        # Each comma-separated range is a separate output PDF (e.g. 1-3,5-7).
        for index, part_spec in enumerate(spec.split(","), start=1):
            part_spec = part_spec.strip()
            if not part_spec:
                continue
            group_pages = _parse_pages(part_spec, total)
            writer = PdfWriter()
            for page_no in group_pages:
                writer.add_page(reader.pages[page_no])
            label = re.sub(r"[^0-9-]+", "", part_spec) or str(index)
            part = outdir / f"{base}-range-{label}.pdf"
            _write_pdf(writer, part); outputs.append((part, part.name))
        if len(outputs) == 1:
            return outputs[0][0], f"Split the requested range into one PDF containing {len(selected)} page(s)."
        return _make_zip(outputs, outdir / f"{base}-split.zip"), f"Split the input into {len(outputs)} PDF range(s)."

    if tool == "compress-pdf":
        target = outdir / f"{base}-compressed.pdf"
        message = _compress_pdf(source, target, str(options.get("level", "recommended")))
        return target, message

    if tool == "organize-pdf":
        order_text = str(options.get("order", "")).strip()
        if order_text:
            ordered_pages = _parse_pages(order_text, total)
        else:
            ordered_pages = list(range(total))
        removed = set(_parse_pages(str(options.get("delete_pages", "")), total)) if str(options.get("delete_pages", "")).strip() else set()
        ordered_pages = [p for p in ordered_pages if p not in removed]
        if not ordered_pages:
            raise HTTPException(400, "Organizing these settings would remove every page.")
        rotations = options.get("rotations", {})
        rotate_selected = set(_parse_pages(str(options.get("rotate_pages", "")), total)) if str(options.get("rotate_pages", "")).strip() else set()
        selected_rotation = int(options.get("rotation", 90))
        if selected_rotation % 90 != 0:
            raise HTTPException(400, "Rotation must be a multiple of 90 degrees.")
        writer = PdfWriter()
        for pno in ordered_pages:
            page = reader.pages[pno]
            mapped_rotation = rotations.get(str(pno + 1), 0) if isinstance(rotations, dict) else 0
            rotation = selected_rotation if pno in rotate_selected else int(mapped_rotation or 0)
            if rotation % 90 != 0:
                raise HTTPException(400, "Rotation must be a multiple of 90 degrees.")
            if rotation:
                page.rotate(rotation)
            writer.add_page(page)
        return _write_pdf(writer, outdir / f"{base}-organized.pdf"), f"Reordered {len(ordered_pages)} page(s), applying requested page rotations and deletions."

    if tool in {"rotate-pdf", "remove-pages", "extract-pages"}:
        if tool == "rotate-pdf":
            selected = _parse_pages(str(options.get("pages", "")), total, default_all=True)
            rotation = int(options.get("rotation", 90))
            if rotation % 90 != 0:
                raise HTTPException(400, "Rotation must be 90, 180, or 270 degrees.")
            writer = PdfWriter()
            selected_set = set(selected)
            for i, page in enumerate(reader.pages):
                if i in selected_set:
                    page.rotate(rotation)
                writer.add_page(page)
            return _write_pdf(writer, outdir / f"{base}-rotated.pdf"), f"Rotated {len(selected)} selected page(s) by {rotation} degrees clockwise."
        selected = _parse_pages(str(options.get("pages", "")), total)
        writer = PdfWriter()
        if tool == "remove-pages":
            selected_set = set(selected)
            for i, page in enumerate(reader.pages):
                if i not in selected_set:
                    writer.add_page(page)
            if len(writer.pages) < 1:
                raise HTTPException(400, "You cannot remove every page from a PDF.")
            return _write_pdf(writer, outdir / f"{base}-pages-removed.pdf"), f"Removed {len(selected)} page(s); {len(writer.pages)} page(s) remain."
        for page_no in selected:
            writer.add_page(reader.pages[page_no])
        return _write_pdf(writer, outdir / f"{base}-extracted.pdf"), f"Extracted {len(selected)} page(s)."

    if tool == "pdf-to-word":
        return outdir / f"{base}.docx", _pdf_to_docx(source, outdir / f"{base}.docx")

    if tool == "pdf-to-jpg":
        dpi = min(240, max(72, int(options.get("dpi", 144))))
        renders = _render_pdf_pages(source, outdir / "jpg-pages", dpi=dpi, ext="jpg")
        return _make_zip([(p, p.name) for p in renders], outdir / f"{base}-jpg-pages.zip"), f"Rendered {len(renders)} page(s) as JPG images at {dpi} DPI."

    if tool == "pdf-to-excel":
        return outdir / f"{base}.xlsx", _pdf_to_xlsx(source, outdir / f"{base}.xlsx")

    if tool == "pdf-to-powerpoint":
        try:
            from pptx import Presentation
            from pptx.util import Inches
            doc = fitz.open(str(source))
            if doc.page_count > 100:
                raise HTTPException(413, "PDF-to-PowerPoint is limited to 100 pages per file.")
            prs = Presentation()
            prs.slide_width = Inches(13.333)
            prs.slide_height = Inches(7.5)
            blank = prs.slide_layouts[6]
            rendered_bytes = 0
            for idx, page in enumerate(doc):
                pix = page.get_pixmap(dpi=150, alpha=False)
                png = outdir / f"ppt-page-{idx + 1:03d}.png"
                pix.save(str(png))
                rendered_bytes += png.stat().st_size
                if rendered_bytes > MAX_OUTPUT_BYTES:
                    doc.close()
                    raise HTTPException(413, f"Rendered slide images exceed the {MAX_OUTPUT_BYTES // (1024 * 1024)} MB output limit.")
                slide = prs.slides.add_slide(blank)
                scale = min(prs.slide_width / pix.width, prs.slide_height / pix.height)
                width = int(pix.width * scale); height = int(pix.height * scale)
                slide.shapes.add_picture(str(png), int((prs.slide_width - width) / 2), int((prs.slide_height - height) / 2), width=width, height=height)
            prs.save(str(outdir / f"{base}.pptx")); doc.close()
            return outdir / f"{base}.pptx", "Created a PowerPoint with one image-based slide per PDF page. Slides preserve page appearance but their text and graphics are not individually editable."
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(422, "Could not convert this PDF to PowerPoint.") from exc

    if tool == "repair-pdf":
        target = outdir / f"{base}-repaired.pdf"
        try:
            doc = fitz.open(str(source))
            doc.save(str(target), garbage=4, deflate=True, clean=True)
            doc.close()
            _read_pdf(target)
            return target, "Rewrote the PDF structure and verified the result can be reopened. This cannot recover missing or destroyed content."
        except Exception:
            gs = shutil.which("gs")
            if gs:
                proc = subprocess.run([gs, "-sDEVICE=pdfwrite", "-dNOPAUSE", "-dQUIET", "-dBATCH", "-dSAFER", f"-sOutputFile={target}", str(source)], capture_output=True, timeout=180, check=False)
                if proc.returncode == 0 and target.exists() and target.stat().st_size:
                    _read_pdf(target)
                    return target, "Rewrote the PDF with Ghostscript and verified the result. Severe corruption or missing content may not be recoverable."
            raise HTTPException(422, "Neither PyMuPDF nor Ghostscript could repair this PDF. The file may be severely damaged or encrypted.")

    if tool == "crop-pdf":
        margins = {k: max(0.0, min(45.0, float(options.get(k, 0)))) for k in ("top", "right", "bottom", "left")}
        try:
            doc = fitz.open(str(source))
            for page in doc:
                w, h = page.rect.width, page.rect.height
                rect = fitz.Rect(w * margins["left"] / 100, h * margins["top"] / 100, w * (1 - margins["right"] / 100), h * (1 - margins["bottom"] / 100))
                if rect.width < 50 or rect.height < 50:
                    raise HTTPException(400, "Crop margins would leave a page too small. Reduce the margins.")
                page.set_cropbox(rect)
            target = outdir / f"{base}-cropped.pdf"
            doc.save(str(target), garbage=4, deflate=True)
            doc.close()
            return target, "Applied crop margins to every page. Cropping changes the visible page area; it does not permanently erase hidden content outside the crop box."
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(422, "Could not crop this PDF.") from exc

    if tool in {"add-text", "page-numbers", "watermark-pdf", "edit-pdf"}:
        try:
            doc = fitz.open(str(source))
            page_list = list(range(doc.page_count))
            if tool in {"add-text", "edit-pdf"} and str(options.get("pages", "")).strip():
                page_list = _parse_pages(str(options.get("pages")), doc.page_count)
            color_in = options.get("color", "#263b59")
            def parse_color(value: Any, default=(0.15, 0.22, 0.33)) -> tuple[float, float, float]:
                if isinstance(value, str) and re.fullmatch(r"#[0-9A-Fa-f]{6}", value):
                    return tuple(int(value[i:i+2], 16) / 255 for i in (1, 3, 5))  # type: ignore
                return default
            color = parse_color(color_in)
            font_size = max(6, min(100, float(options.get("font_size", 18))))
            x_pct = max(0, min(100, float(options.get("x", 10))))
            y_pct = max(0, min(100, float(options.get("y", 20))))
            text = str(options.get("text", ""))[:1000]
            if tool == "page-numbers":
                fmt = str(options.get("format", "{n}"))[:100]
                pos = str(options.get("position", "bottom-center"))
                size = max(6, min(48, float(options.get("font_size", 11))))
                margin = 24
                for i, page in enumerate(doc):
                    label = fmt.replace("{n}", str(i + 1)).replace("{total}", str(doc.page_count))
                    rect = page.rect
                    y = margin if pos.startswith("top") else rect.height - margin - size
                    if pos.endswith("left"):
                        box = fitz.Rect(24, y, rect.width * 0.45, y + size + 8); align = fitz.TEXT_ALIGN_LEFT
                    elif pos.endswith("right"):
                        box = fitz.Rect(rect.width * 0.55, y, rect.width - 24, y + size + 8); align = fitz.TEXT_ALIGN_RIGHT
                    else:
                        box = fitz.Rect(24, y, rect.width - 24, y + size + 8); align = fitz.TEXT_ALIGN_CENTER
                    page.insert_textbox(box, label, fontsize=size, color=color, align=align, fontname="helv")
            elif tool == "watermark-pdf":
                if len(paths) > 1 and paths[1].suffix.lower() in IMAGE_EXTS:
                    with Image.open(paths[1]) as im:
                        rgba = im.convert("RGBA")
                        alpha = rgba.getchannel("A").point(lambda a: int(a * 0.28))
                        rgba.putalpha(alpha)
                        wm_path = outdir / "watermark-alpha.png"; rgba.save(wm_path)
                    for page in doc:
                        r = page.rect
                        w = r.width * 0.52; h = r.height * 0.52
                        rect = fitz.Rect((r.width - w) / 2, (r.height - h) / 2, (r.width + w) / 2, (r.height + h) / 2)
                        page.insert_image(rect, filename=str(wm_path), overlay=True, keep_proportion=True)
                else:
                    if not text.strip():
                        text = "CONFIDENTIAL"
                    for page in doc:
                        r = page.rect
                        box = fitz.Rect(r.width * 0.08, r.height * 0.38, r.width * 0.92, r.height * 0.62)
                        page.insert_textbox(box, text, fontsize=font_size, color=parse_color(options.get("color", "#8995A8"), (0.55, 0.59, 0.66)), align=fitz.TEXT_ALIGN_CENTER, fontname="helv", overlay=True)
            elif tool == "add-text":
                if not text.strip():
                    raise HTTPException(400, "Enter the text you want to add.")
                for index in page_list:
                    page = doc[index]
                    point = fitz.Point(page.rect.width * x_pct / 100, page.rect.height * y_pct / 100)
                    page.insert_text(point, text, fontsize=font_size, color=color, fontname="helv", overlay=True)
            else:  # edit-pdf: text plus one optional shape or image overlay
                for index in page_list:
                    page = doc[index]
                    if text.strip():
                        point = fitz.Point(page.rect.width * x_pct / 100, page.rect.height * y_pct / 100)
                        page.insert_text(point, text, fontsize=font_size, color=color, fontname="helv", overlay=True)
                    shape_name = str(options.get("shape", "none"))
                    left = min(page.rect.width - 2, page.rect.width * x_pct / 100)
                    top = min(page.rect.height - 2, page.rect.height * y_pct / 100)
                    rw = page.rect.width * max(2, min(90, float(options.get("width", 20)))) / 100
                    rh = page.rect.height * max(2, min(90, float(options.get("height", 10)))) / 100
                    rect = fitz.Rect(left, top, min(page.rect.width - 1, left + rw), min(page.rect.height - 1, top + rh))
                    if shape_name == "rectangle": page.draw_rect(rect, color=color, width=2, overlay=True)
                    elif shape_name == "ellipse": page.draw_oval(rect, color=color, width=2, overlay=True)
                    elif shape_name == "line": page.draw_line(rect.tl, rect.br, color=color, width=2, overlay=True)
                    annotation = str(options.get("annotation", "none"))
                    if annotation == "note":
                        note_text = str(options.get("annotation_text", text or "Review note"))[:1000]
                        page.add_text_annot(fitz.Point(left, top), note_text)
                    elif annotation == "highlight":
                        page.add_highlight_annot(rect)
                    elif annotation == "underline":
                        page.add_underline_annot(rect)
                    elif annotation == "strikeout":
                        page.add_strikeout_annot(rect)
                    if len(paths) > 1 and paths[1].suffix.lower() in IMAGE_EXTS:
                        page.insert_image(rect, filename=str(paths[1]), keep_proportion=True, overlay=True)
            target_name = {"add-text": "text-added", "page-numbers": "numbered", "watermark-pdf": "watermarked", "edit-pdf": "edited"}[tool]
            target = outdir / f"{base}-{target_name}.pdf"
            doc.save(str(target), garbage=4, deflate=True)
            doc.close()
            return target, {"add-text": f"Added text to {len(page_list)} page(s). Existing PDF text is not directly rewritten.", "page-numbers": "Added customizable page numbers to every page.", "watermark-pdf": "Applied the selected watermark overlay.", "edit-pdf": f"Applied requested text, shapes, annotations, and/or image overlays on {len(page_list)} page(s). Existing page content is not directly edited."}[tool]
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(422, "The PDF editor could not apply this change.") from exc

    raise HTTPException(501, f"The '{tool}' tool is not implemented.")


async def _cleanup_loop() -> None:
    while True:
        try:
            now = time.time()
            for item in ROOT.iterdir():
                try:
                    if item.is_dir() and now - item.stat().st_mtime > RETENTION_SECONDS:
                        shutil.rmtree(item, ignore_errors=True)
                    elif item.is_file() and now - item.stat().st_mtime > RETENTION_SECONDS:
                        item.unlink(missing_ok=True)
                except OSError:
                    continue
        except Exception:
            pass
        await asyncio.sleep(60)


@asynccontextmanager
async def lifespan(_: FastAPI):
    task = asyncio.create_task(_cleanup_loop())
    yield
    task.cancel()
    try:
        await task
    except asyncio.CancelledError:
        pass


app = FastAPI(title=APP_NAME, version="1.0.0", description="Temporary-file PDF processing API", lifespan=lifespan)
app.add_middleware(CORSMiddleware, allow_origins=ALLOWED_ORIGINS, allow_credentials=False, allow_methods=["GET", "POST", "OPTIONS"], allow_headers=["*"])


@app.middleware("http")
async def add_private_document_headers(request, call_next):
    response = await call_next(request)
    response.headers.setdefault("Cache-Control", "no-store")
    response.headers.setdefault("X-Content-Type-Options", "nosniff")
    response.headers.setdefault("Referrer-Policy", "no-referrer")
    return response


@app.get("/api/health")
def health() -> dict[str, Any]:
    return {"status": "ok", "service": APP_NAME, "max_file_bytes": MAX_FILE_BYTES, "max_files": MAX_FILES, "retention_seconds": RETENTION_SECONDS}


@app.post("/api/preview")
async def preview(file: UploadFile = File(...)) -> dict[str, Any]:
    """Transient PDF preview: only low-resolution thumbnails and page count are returned; the source is not persisted."""
    tmpdir = Path(tempfile.mkdtemp(prefix="preview-", dir=str(ROOT)))
    path = tmpdir / (Path(file.filename or "document.pdf").name)
    try:
        if path.suffix.lower() != ".pdf":
            raise HTTPException(400, "Preview thumbnails are currently available for PDF files.")
        chunks = 0; first = b""
        with path.open("wb") as handle:
            while True:
                chunk = await file.read(1024 * 1024)
                if not chunk: break
                chunks += len(chunk)
                if chunks > MAX_FILE_BYTES:
                    raise HTTPException(413, f"File exceeds the {MAX_FILE_BYTES // (1024 * 1024)} MB upload limit.")
                if len(first) < 1024: first += chunk[:1024 - len(first)]
                handle.write(chunk)
        _check_file(path, first)
        doc = fitz.open(str(path))
        if doc.page_count < 1: raise HTTPException(422, "This PDF has no pages.")
        if doc.page_count > MAX_PDF_PAGES: raise HTTPException(413, f"PDF exceeds the {MAX_PDF_PAGES}-page limit.")
        page_count = doc.page_count
        step = max(1, (page_count + 5) // 6)
        page_dimensions = [{"page": index + 1, "width_pt": float(page.rect.width), "height_pt": float(page.rect.height)} for index, page in enumerate(doc)]
        thumbs = []
        for i in range(0, doc.page_count, step):
            pix = doc[i].get_pixmap(matrix=fitz.Matrix(0.45, 0.45), alpha=False)
            encoded = base64.b64encode(pix.tobytes("png")).decode("ascii")
            thumbs.append({"page": i + 1, "data_url": f"data:image/png;base64,{encoded}"})
            if len(thumbs) >= 6: break
        doc.close()
        return {"filename": path.name, "page_count": page_count, "page_dimensions": page_dimensions, "thumbnails": thumbs}
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, "Could not preview this PDF. The file may be corrupted, encrypted, or unsupported.") from exc
    finally:
        await file.close()
        shutil.rmtree(tmpdir, ignore_errors=True)


@app.post("/api/organize-preview")
async def organize_preview(file: UploadFile = File(...)) -> dict[str, Any]:
    """Render a compact thumbnail for every page to support visual drag-and-drop organization."""
    tmpdir = Path(tempfile.mkdtemp(prefix="organize-preview-", dir=str(ROOT)))
    path = tmpdir / Path(file.filename or "document.pdf").name
    try:
        if path.suffix.lower() != ".pdf":
            raise HTTPException(400, "Organize PDF preview requires a PDF file.")
        total_bytes = 0
        first = b""
        with path.open("wb") as handle:
            while True:
                chunk = await file.read(1024 * 1024)
                if not chunk:
                    break
                total_bytes += len(chunk)
                if total_bytes > MAX_FILE_BYTES:
                    raise HTTPException(413, f"File exceeds the {MAX_FILE_BYTES // (1024 * 1024)} MB upload limit.")
                if len(first) < 1024:
                    first += chunk[:1024 - len(first)]
                handle.write(chunk)
        _check_file(path, first)
        doc = fitz.open(str(path))
        if doc.is_encrypted:
            doc.close()
            raise HTTPException(422, "This PDF is password-protected. Unlock it before organizing pages.")
        if doc.page_count < 1:
            doc.close()
            raise HTTPException(422, "This PDF has no pages.")
        if doc.page_count > MAX_PDF_PAGES:
            doc.close()
            raise HTTPException(413, f"Visual organization is limited to {MAX_PDF_PAGES} pages per PDF.")
        pages = []
        for index, page in enumerate(doc, start=1):
            rect = page.rect
            scale = min(0.22, 140 / max(rect.width, rect.height))
            pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False, annots=False)
            image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            buffer = io.BytesIO()
            image.save(buffer, format="JPEG", quality=58, optimize=True)
            encoded = base64.b64encode(buffer.getvalue()).decode("ascii")
            pages.append({"page": index, "width_pt": float(rect.width), "height_pt": float(rect.height), "data_url": f"data:image/jpeg;base64,{encoded}"})
        page_count = doc.page_count
        doc.close()
        return {"filename": path.name, "page_count": page_count, "pages": pages, "max_pages": MAX_PDF_PAGES}
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, "Could not create page thumbnails. The PDF may be corrupted, encrypted, or unsupported.") from exc
    finally:
        await file.close()
        shutil.rmtree(tmpdir, ignore_errors=True)


@app.post("/api/redact-preview")
async def create_redact_preview(file: UploadFile = File(...)) -> dict[str, Any]:
    """Create a short-lived preview session so the user can draw redaction regions on page images."""
    session_id = uuid.uuid4().hex
    session_dir = ROOT / f"redact-preview-{session_id}"
    session_dir.mkdir(parents=True, mode=0o700, exist_ok=False)
    path = session_dir / "source.pdf"
    try:
        if Path(file.filename or "document.pdf").suffix.lower() != ".pdf":
            raise HTTPException(400, "Redaction preview requires a PDF file.")
        total = 0
        first = b""
        with path.open("wb") as handle:
            while True:
                chunk = await file.read(1024 * 1024)
                if not chunk:
                    break
                total += len(chunk)
                if total > MAX_FILE_BYTES:
                    raise HTTPException(413, f"File exceeds the {MAX_FILE_BYTES // (1024 * 1024)} MB upload limit.")
                if len(first) < 1024:
                    first += chunk[:1024 - len(first)]
                handle.write(chunk)
        path.chmod(0o600)
        _check_file(path, first)
        with fitz.open(str(path)) as doc:
            if doc.needs_pass:
                raise HTTPException(422, "This PDF is password-protected. Unlock it before selecting redaction regions.")
            if doc.page_count < 1:
                raise HTTPException(422, "This PDF has no pages.")
            if doc.page_count > MAX_PDF_PAGES:
                raise HTTPException(413, f"Redaction preview is limited to {MAX_PDF_PAGES} pages.")
            page_count = doc.page_count
        return {"session_id": session_id, "page_count": page_count, "retention_seconds": RETENTION_SECONDS}
    except HTTPException:
        shutil.rmtree(session_dir, ignore_errors=True)
        raise
    except Exception as exc:
        shutil.rmtree(session_dir, ignore_errors=True)
        raise HTTPException(422, "Could not open a redaction preview for this PDF.") from exc
    finally:
        await file.close()


@app.get("/api/redact-preview/{session_id}/page/{page_number}")
def redact_preview_page(session_id: str, page_number: int) -> dict[str, Any]:
    """Render one page from a temporary redaction session in page-point coordinates."""
    if not re.fullmatch(r"[a-f0-9]{32}", session_id):
        raise HTTPException(404, "Redaction preview has expired. Upload the PDF again.")
    session_dir = ROOT / f"redact-preview-{session_id}"
    path = session_dir / "source.pdf"
    if not path.is_file():
        raise HTTPException(404, "Redaction preview has expired. Upload the PDF again.")
    try:
        with fitz.open(str(path)) as doc:
            if doc.needs_pass:
                raise HTTPException(422, "This PDF is password-protected. Unlock it before selecting redaction regions.")
            if page_number < 1 or page_number > doc.page_count:
                raise HTTPException(404, f"Page {page_number} does not exist in this PDF.")
            page = doc[page_number - 1]
            rect = page.rect
            scale = min(1.4, 1600 / max(rect.width, rect.height))
            pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False, annots=True)
            encoded = base64.b64encode(pix.tobytes("png")).decode("ascii")
            response = {
                "page": page_number,
                "page_count": doc.page_count,
                "width_pt": float(rect.width),
                "height_pt": float(rect.height),
                "pixel_width": pix.width,
                "pixel_height": pix.height,
                "data_url": f"data:image/png;base64,{encoded}",
            }
        # Refresh retention while the user is actively reviewing pages.
        os.utime(session_dir, None)
        return response
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, "Could not render this PDF page for redaction selection.") from exc


@app.post("/api/form-fields")
async def inspect_form_fields(file: UploadFile = File(...)) -> dict[str, Any]:
    """Inspect supported AcroForm fields from a temporary upload; no source document is retained."""
    tmpdir = Path(tempfile.mkdtemp(prefix="form-inspect-", dir=str(ROOT)))
    path = tmpdir / Path(file.filename or "form.pdf").name
    try:
        if path.suffix.lower() != ".pdf":
            raise HTTPException(400, "Upload a PDF file to inspect its form fields.")
        total = 0
        first = b""
        with path.open("wb") as handle:
            while True:
                chunk = await file.read(1024 * 1024)
                if not chunk:
                    break
                total += len(chunk)
                if total > MAX_FILE_BYTES:
                    raise HTTPException(413, f"File exceeds the {MAX_FILE_BYTES // (1024 * 1024)} MB upload limit.")
                if len(first) < 1024:
                    first += chunk[:1024 - len(first)]
                handle.write(chunk)
        _check_file(path, first)
        fields = await run_in_threadpool(_form_field_details, path)
        return {"filename": path.name, "field_count": len(fields), "fields": fields, "fillable": bool(fields)}
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, "Could not inspect this PDF form.") from exc
    finally:
        await file.close()
        shutil.rmtree(tmpdir, ignore_errors=True)


@app.post("/api/process/{tool}")
async def process(tool: str, files: list[UploadFile] = File(default=[]), options: str = Form("{}")) -> dict[str, Any]:
    if tool not in PDF_TOOLS:
        raise HTTPException(404, "Unknown tool.")
    opts = _parse_options(options)
    if not files and not (tool == "html-to-pdf" and str(opts.get("url", "")).strip()):
        raise HTTPException(400, "Upload at least one file, or enter a public URL for HTML-to-PDF.")
    if len(files) > MAX_FILES:
        raise HTTPException(413, f"You can upload at most {MAX_FILES} files per operation.")
    job_id = uuid.uuid4().hex
    job_dir = ROOT / job_id
    job_dir.mkdir(parents=True, exist_ok=False, mode=0o700)
    source_dir = job_dir / "source"; source_dir.mkdir(mode=0o700)
    work_dir = job_dir / "work"; work_dir.mkdir(mode=0o700)
    saved: list[Path] = []
    total_bytes = 0
    try:
        for index, upload in enumerate(files, start=1):
            original = Path(upload.filename or f"file-{index}").name
            ext = Path(original).suffix.lower()
            if ext not in ALL_EXTS:
                raise HTTPException(400, f"Unsupported file type for upload {index}. Supported types: PDF, DOC/DOCX, XLS/XLSX, PPT/PPTX, JPG, PNG, and HTML for HTML-to-PDF.")
            dest = source_dir / f"{index:02d}-{safe_stem(original)}{ext}"
            size = 0; first = b""
            with dest.open("wb") as output:
                while True:
                    chunk = await upload.read(1024 * 1024)
                    if not chunk: break
                    size += len(chunk); total_bytes += len(chunk)
                    if size > MAX_FILE_BYTES:
                        raise HTTPException(413, f"Each file is limited to {MAX_FILE_BYTES // (1024 * 1024)} MB.")
                    if total_bytes > MAX_TOTAL_BYTES:
                        raise HTTPException(413, f"Total upload is limited to {MAX_TOTAL_BYTES // (1024 * 1024)} MB per operation.")
                    if len(first) < 1024: first += chunk[:1024 - len(first)]
                    output.write(chunk)
            try:
                dest.chmod(0o600)
            except OSError:
                pass
            _check_file(dest, first)
            saved.append(dest)

        labels = [safe_stem(Path(upload.filename or f"file-{i}").name) for i, upload in enumerate(files, start=1)]
        if labels:
            base = labels[0] if len(set(labels)) == 1 or len(labels) == 1 else "papernest-output"
        elif tool == "html-to-pdf" and str(opts.get("url", "")).strip():
            base = safe_stem(urlsplit(str(opts["url"])).hostname or "webpage")
        else:
            base = "papernest-output"
        message = ""
        if tool in PDF_BATCHABLE and len(saved) > 1:
            # A second image is an optional watermark/editor asset, not another document in batch mode.
            image_asset = tool in {"watermark-pdf", "edit-pdf"} and len(saved) == 2 and saved[1].suffix.lower() in IMAGE_EXTS
            if image_asset:
                result, message = await run_in_threadpool(_process_single, tool, saved, opts, work_dir / "result", base)
            else:
                artifacts: list[tuple[Path, str]] = []
                messages = []
                for index, path in enumerate(saved, start=1):
                    local_dir = work_dir / f"batch-{index}"
                    artifact, item_message = await run_in_threadpool(_process_single, tool, [path], opts, local_dir, safe_stem(path.name))
                    messages.append(item_message)
                    if artifact.suffix.lower() == ".zip":
                        try:
                            with zipfile.ZipFile(artifact) as zf:
                                for member in zf.namelist():
                                    if member.endswith("/"): continue
                                    # Flatten nested output zips and avoid unsafe archive paths.
                                    flat_name = f"{safe_stem(path.name)}-{Path(member).name}"
                                    temp_out = work_dir / f"flattened-{uuid.uuid4().hex}-{Path(member).name}"
                                    temp_out.write_bytes(zf.read(member))
                                    artifacts.append((temp_out, flat_name))
                        except zipfile.BadZipFile:
                            artifacts.append((artifact, artifact.name))
                    else:
                        artifacts.append((artifact, f"{safe_stem(path.name)}-{artifact.name}"))
                result = await run_in_threadpool(_make_zip, artifacts, work_dir / "batch-results.zip")
                message = f"Processed {len(saved)} files and bundled the results. " + " ".join(dict.fromkeys(m for m in messages if m))
        else:
            result, message = await run_in_threadpool(_process_single, tool, saved, opts, work_dir / "result", base)

        if not result.exists() or result.stat().st_size == 0:
            raise HTTPException(500, "The processor did not create an output file. Please retry with a different file.")
        if result.stat().st_size > MAX_OUTPUT_BYTES:
            raise HTTPException(413, f"Generated output exceeds the {MAX_OUTPUT_BYTES // (1024 * 1024)} MB output limit.")
        if result.suffix.lower() == ".pdf":
            if tool == "protect-pdf":
                protected_reader = PdfReader(str(result), strict=False)
                if not protected_reader.is_encrypted:
                    raise HTTPException(500, "The protected output did not retain encryption.")
            else:
                await run_in_threadpool(_read_pdf, result)
        # Move final output to a flat, unguessable job path; download endpoint never accepts caller-supplied paths.
        output_name = result.name
        final_path = job_dir / f"output{result.suffix.lower()}"
        shutil.copyfile(result, final_path)
        try:
            final_path.chmod(0o600)
        except OSError:
            pass
        (job_dir / "metadata.json").write_text(json.dumps({"name": output_name, "size": final_path.stat().st_size, "created": time.time()}), encoding="utf-8")
        return {"job_id": job_id, "filename": output_name, "size": final_path.stat().st_size, "download_url": f"/api/download/{job_id}", "message": message or "Processing completed successfully."}
    except HTTPException:
        shutil.rmtree(job_dir, ignore_errors=True)
        raise
    except Exception:
        shutil.rmtree(job_dir, ignore_errors=True)
        raise HTTPException(500, "Unexpected processing error. No output was retained. Please retry or use a different file.")
    finally:
        for upload in files:
            await upload.close()


@app.get("/api/download/{job_id}")
def download(job_id: str) -> FileResponse:
    if not re.fullmatch(r"[a-f0-9]{32}", job_id):
        raise HTTPException(404, "Download not found or expired.")
    job_dir = ROOT / job_id
    metadata_path = job_dir / "metadata.json"
    output_candidates = list(job_dir.glob("output.*")) if job_dir.exists() else []
    if not metadata_path.exists() or len(output_candidates) != 1:
        raise HTTPException(404, "Download not found or expired. Temporary files are deleted automatically after 30 minutes.")
    try:
        metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
    except Exception as exc:
        raise HTTPException(404, "Download is no longer available.") from exc
    output_path = output_candidates[0]
    name = re.sub(r"[^A-Za-z0-9._-]+", "_", str(metadata.get("name", output_path.name)))[:120]
    return FileResponse(output_path, filename=name, media_type=MIME_TYPES.get(output_path.suffix.lower(), "application/octet-stream"), headers={"Cache-Control": "no-store", "X-Content-Type-Options": "nosniff"})
